"""Build the Amazon fee-invoice register workbook from parsed documents."""

import sys
import collections
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from amazon_fee_invoice import parse

INK, CLAY, CREAM = '1E3A57', 'B4633A', 'FAF6F1'
OK_BG, OK_FG, BAD_BG, BAD_FG, HILIGHT = 'E7F1EB', '256B4A', 'FBEBE9', '9B2F2B', 'EAF0F6'
RULE = 'D8DEE6'

FMT_AMT = '#,##0.00;-#,##0.00'
FMT_INT = '#,##0'

STATES = {'01': 'Jammu & Kashmir', '02': 'Himachal Pradesh', '03': 'Punjab', '04': 'Chandigarh',
          '05': 'Uttarakhand', '06': 'Haryana', '07': 'Delhi', '08': 'Rajasthan',
          '09': 'Uttar Pradesh', '10': 'Bihar', '11': 'Sikkim', '12': 'Arunachal Pradesh',
          '13': 'Nagaland', '14': 'Manipur', '15': 'Mizoram', '16': 'Tripura',
          '17': 'Meghalaya', '18': 'Assam', '19': 'West Bengal', '20': 'Jharkhand',
          '21': 'Odisha', '22': 'Chhattisgarh', '23': 'Madhya Pradesh', '24': 'Gujarat',
          '27': 'Maharashtra', '29': 'Karnataka', '30': 'Goa', '32': 'Kerala',
          '33': 'Tamil Nadu', '34': 'Puducherry', '36': 'Telangana', '37': 'Andhra Pradesh'}

thin = Side(style='thin', color=RULE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def head(ws, row, cols, labels):
    for c, label in zip(cols, labels):
        cell = ws.cell(row=row, column=c, value=label)
        cell.font = Font(bold=True, color='FFFFFF', size=10)
        cell.fill = PatternFill('solid', fgColor=INK)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BOX
    ws.row_dimensions[row].height = 28


def put(ws, r, c, v, *, fmt=None, bold=False, fill=None, fg=None, align=None, border=True):
    cell = ws.cell(row=r, column=c, value=v)
    if fmt:
        cell.number_format = fmt
    cell.font = Font(bold=bold, size=10, color=fg or '1F2933')
    if fill:
        cell.fill = PatternFill('solid', fgColor=fill)
    if align:
        cell.alignment = Alignment(horizontal=align)
    if border:
        cell.border = BOX
    return cell


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def tax_of(row, head_name):
    return sum((t['amount'] for t in row['taxes'] if t['tax'] == head_name), Decimal(0))


def build(docs, out_path, period_label):
    wb = Workbook()

    # ---------------------------------------------------------------- Summary
    ws = wb.active
    ws.title = 'Summary'
    ws.sheet_view.showGridLines = False
    widths(ws, {'A': 3, 'B': 34, 'C': 10, 'D': 9, 'E': 14, 'F': 12, 'G': 12, 'H': 13, 'I': 14})

    ws.merge_cells('B2:I2')
    t = ws['B2']
    t.value = 'Amazon Fee Invoices — Input Credit Register'
    t.font = Font(bold=True, size=16, color=INK)
    ws.merge_cells('B3:I3')
    ws['B3'].value = f'{docs[0]["header"].get("buyer_name", "")}   ·   {period_label}'
    ws['B3'].font = Font(size=10, color=CLAY)
    ws.row_dimensions[2].height = 24

    invs = [d for d in docs if d['header']['doc_type'] == 'INVOICE']
    cns = [d for d in docs if d['header']['doc_type'] == 'CREDIT_NOTE']
    r = 5
    for label, val in (('Tax invoices', len(invs)), ('Credit notes', len(cns)),
                       ('Fee lines', sum(len(d['detail']) for d in docs))):
        put(ws, r, 2, label, bold=True, fill=CREAM)
        put(ws, r, 3, val, fmt=FMT_INT, align='right', fill=CREAM)
        r += 1

    # ----- fee type breakdown
    r += 1
    put(ws, r, 2, 'WHAT AMAZON CHARGED', bold=True, fg=CLAY, border=False)
    r += 1
    hrow = r
    head(ws, r, range(2, 10), ['Fee type', 'SAC', 'Lines', 'Taxable value',
                               'CGST', 'SGST', 'IGST', 'Total'])
    r += 1

    agg = collections.defaultdict(lambda: {'n': 0, 'tax': Decimal(0), 'CGST': Decimal(0),
                                           'SGST': Decimal(0), 'IGST': Decimal(0), 'sac': ''})
    for d in docs:
        for row in d['detail']:
            a = agg[row['description']]
            a['n'] += 1
            a['sac'] = row['sac']
            a['tax'] += row['amount']
            for hd in ('CGST', 'SGST', 'IGST'):
                a[hd] += tax_of(row, hd)

    first = r
    for name, a in sorted(agg.items(), key=lambda x: -x[1]['tax']):
        put(ws, r, 2, name)
        put(ws, r, 3, a['sac'], align='center')
        put(ws, r, 4, a['n'], fmt=FMT_INT, align='right')
        put(ws, r, 5, float(a['tax']), fmt=FMT_AMT)
        put(ws, r, 6, float(a['CGST']), fmt=FMT_AMT)
        put(ws, r, 7, float(a['SGST']), fmt=FMT_AMT)
        put(ws, r, 8, float(a['IGST']), fmt=FMT_AMT)
        put(ws, r, 9, f'=SUM(E{r}:H{r})', fmt=FMT_AMT)
        r += 1
    last = r - 1
    put(ws, r, 2, 'Total', bold=True, fill=HILIGHT)
    put(ws, r, 3, '', fill=HILIGHT)
    put(ws, r, 4, f'=SUM(D{first}:D{last})', fmt=FMT_INT, bold=True, fill=HILIGHT, align='right')
    for c in range(5, 10):
        L = get_column_letter(c)
        put(ws, r, c, f'=SUM({L}{first}:{L}{last})', fmt=FMT_AMT, bold=True, fill=HILIGHT)
    total_row = r
    ws.auto_filter.ref = f'B{hrow}:I{last}'

    # ----- by month
    r += 2
    put(ws, r, 2, 'BY MONTH', bold=True, fg=CLAY, border=False)
    r += 1
    head(ws, r, range(2, 7), ['Month', 'Documents', 'Taxable value', 'Input credit', 'Total'])
    r += 1
    permonth = collections.defaultdict(lambda: {'n': 0, 'tax': Decimal(0), 'gst': Decimal(0)})
    for d in docs:
        dt = d['header'].get('doc_date', '')          # DD/MM/YYYY
        key = (dt[6:], dt[3:5])
        permonth[key]['n'] += 1
        permonth[key]['tax'] += d['calculated']['taxable']
        permonth[key]['gst'] += d['calculated']['gst']
    MONTHS = ['', 'January', 'February', 'March', 'April', 'May', 'June', 'July',
              'August', 'September', 'October', 'November', 'December']
    mfirst = r
    for key in sorted(permonth):
        a = permonth[key]
        put(ws, r, 2, f'{MONTHS[int(key[1])]} {key[0]}')
        put(ws, r, 3, a['n'], fmt=FMT_INT, align='right')
        put(ws, r, 4, float(a['tax']), fmt=FMT_AMT)
        put(ws, r, 5, float(a['gst']), fmt=FMT_AMT)
        put(ws, r, 6, f'=D{r}+E{r}', fmt=FMT_AMT)
        r += 1
    mlast = r - 1
    put(ws, r, 2, 'Total', bold=True, fill=HILIGHT)
    for c in range(3, 7):
        L = get_column_letter(c)
        put(ws, r, c, f'=SUM({L}{mfirst}:{L}{mlast})',
            fmt=FMT_INT if c == 3 else FMT_AMT, bold=True, fill=HILIGHT,
            align='right' if c == 3 else None)

    # ----- ITC by registration
    r += 2
    put(ws, r, 2, 'INPUT CREDIT — BY GST REGISTRATION', bold=True, fg=CLAY, border=False)
    r += 1
    put(ws, r, 2, 'Each state\'s credit is claimed only in that state\'s own GSTR-3B.',
        fg='5B6B7B', border=False)
    r += 1
    head(ws, r, range(2, 8), ['GSTIN (Koparo)', 'State', 'Documents',
                              'Taxable value', 'Input credit', 'Total'])
    r += 1
    per = collections.defaultdict(lambda: {'n': 0, 'tax': Decimal(0), 'gst': Decimal(0)})
    for d in docs:
        g = d['header'].get('buyer_gstin', '?')
        per[g]['n'] += 1
        per[g]['tax'] += d['calculated']['taxable']
        per[g]['gst'] += d['calculated']['gst']
    gfirst = r
    for g, a in sorted(per.items(), key=lambda x: -x[1]['gst']):
        put(ws, r, 2, g)
        put(ws, r, 3, STATES.get(g[:2], g[:2]))
        put(ws, r, 4, a['n'], fmt=FMT_INT, align='right')
        put(ws, r, 5, float(a['tax']), fmt=FMT_AMT)
        put(ws, r, 6, float(a['gst']), fmt=FMT_AMT)
        put(ws, r, 7, f'=E{r}+F{r}', fmt=FMT_AMT)
        r += 1
    glast = r - 1
    put(ws, r, 2, 'Total', bold=True, fill=HILIGHT)
    put(ws, r, 3, '', fill=HILIGHT)
    for c in range(4, 8):
        L = get_column_letter(c)
        put(ws, r, c, f'=SUM({L}{gfirst}:{L}{glast})',
            fmt=FMT_INT if c == 4 else FMT_AMT, bold=True, fill=HILIGHT,
            align='right' if c == 4 else None)
    gtotal = r

    # ----- the two ties that make the register trustworthy
    r += 2
    put(ws, r, 2, 'CHECKS', bold=True, fg=CLAY, border=False)
    r += 1
    head(ws, r, range(2, 5), ['Check', 'Result', 'Must be'])
    r += 1
    put(ws, r, 2, 'Fee-type table equals registration table (taxable)')
    put(ws, r, 3, f'=ROUND(E{total_row}-E{gtotal},2)', fmt=FMT_AMT)
    put(ws, r, 4, 0, fmt=FMT_AMT)
    r += 1
    put(ws, r, 2, 'Fee-type table equals registration table (tax)')
    put(ws, r, 3, f'=ROUND(F{total_row}+G{total_row}+H{total_row}-F{gtotal},2)', fmt=FMT_AMT)
    put(ws, r, 4, 0, fmt=FMT_AMT)
    r += 1
    npass = sum(1 for d in docs for v in d['checks'].values() if v)
    ntot = sum(len(d['checks']) for d in docs)
    put(ws, r, 2, 'Every document ties to the totals Amazon printed on it')
    c = put(ws, r, 3, f'{npass} of {ntot}', align='center')
    c.fill = PatternFill('solid', fgColor=OK_BG if npass == ntot else BAD_BG)
    c.font = Font(bold=True, size=10, color=OK_FG if npass == ntot else BAD_FG)
    put(ws, r, 4, 'all', align='center')

    # ------------------------------------------------------------- Documents
    ws2 = wb.create_sheet('Documents')
    ws2.sheet_view.showGridLines = False
    widths(ws2, {'A': 3, 'B': 20, 'C': 12, 'D': 13, 'E': 18, 'F': 16, 'G': 18,
                 'H': 14, 'I': 12, 'J': 12, 'K': 12, 'L': 14, 'M': 10})
    head(ws2, 2, range(2, 14), ['Document no', 'Date', 'Type', 'Supplier GSTIN',
                                'Place of supply', 'Our GSTIN', 'Taxable value',
                                'CGST', 'SGST', 'IGST', 'Document total', 'Ties'])
    rr = 3
    for d in sorted(docs, key=lambda x: (x['header'].get('doc_date', '')[6:],
                                         x['header'].get('doc_date', '')[3:5],
                                         x['header'].get('doc_no', ''))):
        h = d['header']
        bh = d['calculated']['by_head']
        put(ws2, rr, 2, h.get('doc_no'))
        put(ws2, rr, 3, h.get('doc_date'), align='center')
        put(ws2, rr, 4, 'Credit note' if h['doc_type'] == 'CREDIT_NOTE' else 'Tax invoice')
        put(ws2, rr, 5, h.get('seller_gstin'))
        put(ws2, rr, 6, h.get('place_of_supply'))
        put(ws2, rr, 7, h.get('buyer_gstin'))
        put(ws2, rr, 8, float(d['calculated']['taxable']), fmt=FMT_AMT)
        put(ws2, rr, 9, float(bh.get('CGST', 0)), fmt=FMT_AMT)
        put(ws2, rr, 10, float(bh.get('SGST', 0)), fmt=FMT_AMT)
        put(ws2, rr, 11, float(bh.get('IGST', 0)), fmt=FMT_AMT)
        put(ws2, rr, 12, f'=SUM(H{rr}:K{rr})', fmt=FMT_AMT, bold=True)
        ok = d['ok']
        c = put(ws2, rr, 13, 'Yes' if ok else 'No', align='center')
        c.fill = PatternFill('solid', fgColor=OK_BG if ok else BAD_BG)
        c.font = Font(bold=True, size=10, color=OK_FG if ok else BAD_FG)
        rr += 1
    put(ws2, rr, 2, 'Total', bold=True, fill=HILIGHT)
    for c in range(3, 8):
        put(ws2, rr, c, '', fill=HILIGHT)
    for c in range(8, 13):
        L = get_column_letter(c)
        put(ws2, rr, c, f'=SUM({L}3:{L}{rr-1})', fmt=FMT_AMT, bold=True, fill=HILIGHT)
    put(ws2, rr, 13, '', fill=HILIGHT)
    ws2.auto_filter.ref = f'B2:M{rr-1}'
    ws2.freeze_panes = 'B3'

    # ------------------------------------------------------------- Fee lines
    ws3 = wb.create_sheet('Fee Lines')
    ws3.sheet_view.showGridLines = False
    widths(ws3, {'A': 3, 'B': 20, 'C': 12, 'D': 9, 'E': 32, 'F': 14, 'G': 8,
                 'H': 7, 'I': 12, 'J': 14, 'K': 20})
    head(ws3, 2, range(2, 12), ['Document no', 'Fee date', 'SAC', 'Fee type',
                                'Taxable value', 'Tax', 'Rate', 'Tax amount',
                                'Line total', 'Credit against'])
    rr = 3
    for d in docs:
        for row in d['detail']:
            order = {'CGST': 0, 'SGST': 1, 'UTGST': 2, 'IGST': 3, 'CESS': 4}
            heads = sorted(row['taxes'], key=lambda t: order.get(t['tax'], 9))
            put(ws3, rr, 2, d['header'].get('doc_no'))
            put(ws3, rr, 3, row['fee_date'], align='center')
            put(ws3, rr, 4, row['sac'], align='center')
            put(ws3, rr, 5, row['description'])
            put(ws3, rr, 6, float(row['amount']), fmt=FMT_AMT)
            put(ws3, rr, 7, '+'.join(t['tax'] for t in heads) or '-', align='center')
            # the combined rate, so it reads against the combined tax amount
            # beside it -- CGST 9% + SGST 9% is an 18% line, not a 9% one
            rate = sum((t['rate'] for t in heads), Decimal(0))
            put(ws3, rr, 8, float(rate) / 100 if heads else None,
                fmt='0.00%', align='center')
            put(ws3, rr, 9, float(sum(t['amount'] for t in row['taxes'])), fmt=FMT_AMT)
            put(ws3, rr, 10, f'=F{rr}+I{rr}', fmt=FMT_AMT)
            put(ws3, rr, 11, row['original_invoice_no'] or '')
            rr += 1
    put(ws3, rr, 2, 'Total', bold=True, fill=HILIGHT)
    for c in (3, 4, 5, 7, 8, 11):
        put(ws3, rr, c, '', fill=HILIGHT)
    for c in (6, 9, 10):
        L = get_column_letter(c)
        put(ws3, rr, c, f'=SUM({L}3:{L}{rr-1})', fmt=FMT_AMT, bold=True, fill=HILIGHT)
    ws3.auto_filter.ref = f'B2:K{rr-1}'
    ws3.freeze_panes = 'B3'

    wb.save(out_path)
    return out_path


if __name__ == '__main__':
    out = sys.argv[1]
    period = sys.argv[2]
    docs = [parse(p) for p in sys.argv[3:]]
    print('written:', build(docs, out, period))
