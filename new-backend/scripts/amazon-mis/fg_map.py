"""
SKU -> the brand's own Tally stock name ("FG"), and SKU -> HSN, lifted from the
accountant's existing working files.

The MTR gives Amazon's SKU and Amazon's long marketing description. The
accountant's macro adds an FG column carrying the name actually used in the
books. Showing that name means he reads the COGS sheet in his own vocabulary
rather than translating from Amazon's.

The same files also carry HSN, so they are read for it too — and they confirm
the gap is real rather than a parsing failure: the eight SKUs missing an HSN in
the MTR are missing it here as well, and flow into GSTR-1 blank today.
"""

import glob
import os
from openpyxl import load_workbook


def _cols(ws):
    return [str(c.value).strip() if c.value else '' for c in ws[1]]


def load(workings_dir):
    fg, hsn = {}, {}
    for path in sorted(glob.glob(os.path.join(workings_dir, '*.xlsx'))):
        try:
            wb = load_workbook(path, data_only=True, read_only=True)
        except Exception:
            continue
        for name in wb.sheetnames:
            if 'raw' not in name.lower():
                continue
            ws = wb[name]
            hdr = _cols(ws)
            if 'Sku' not in hdr:
                continue
            i_sku = hdr.index('Sku')
            i_fg = hdr.index('FG') if 'FG' in hdr else None
            i_hsn = hdr.index('Hsn/sac') if 'Hsn/sac' in hdr else None
            for row in ws.iter_rows(min_row=2, values_only=True):
                if i_sku >= len(row) or not row[i_sku]:
                    continue
                sku = str(row[i_sku]).strip()
                if i_fg is not None and i_fg < len(row) and row[i_fg]:
                    fg.setdefault(sku, str(row[i_fg]).strip())
                if i_hsn is not None and i_hsn < len(row) and row[i_hsn]:
                    hsn.setdefault(sku, str(row[i_hsn]).strip())
        wb.close()
    return fg, hsn


if __name__ == '__main__':
    import sys
    f, h = load(sys.argv[1])
    print(f'FG names: {len(f)}   HSN: {len(h)}')
