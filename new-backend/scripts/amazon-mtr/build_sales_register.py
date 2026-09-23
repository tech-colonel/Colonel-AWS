"""
Amazon MTR -> monthly sales register workbook.

The point of this book is that a preparer can file GSTR-1 from it without
reopening the CSVs, and that a reviewer can see, on the same page, what was NOT
proved.  So the failures get a sheet of their own and the Summary carries a
loud pointer to it.  A register that quietly shows only the rows that footed is
worse than no register.

Sheets
  Summary          per month, B2B vs B2C: gross / returns / net taxable, tax by
                   head, TCS.  Stock transfers sit in their own block because
                   they are not turnover.
  State-wise       per month, origin state x destination state.  Intra-state
                   must be CGST+SGST (or UTGST), inter-state must be IGST; the
                   sheet says which combination each pair used and flags any
                   pair that broke the rule.
  Detail           every parsed row.
  Stock Transfers  FC-to-FC movements on Koparo's own GSTINs.
  Exceptions       every verification failure, with file, line and amounts.

Two conventions this book depends on, both proved in mtr_parser.verify_row():

  * Refunds are ALREADY negative in Amazon's file.  Nothing here flips a sign;
    "returns" is simply the sum of rows whose transaction type is Refund, and
    net = gross + returns because returns arrive negative.

  * A refund is dated by its CREDIT NOTE date, not its invoice date.  147 of
    147 refunds across these five months carry an invoice date from an earlier
    month, so grouping on invoice date would post returns into months already
    filed.

Money is formatted '#,##0.00;-#,##0.00' so a negative reads as a negative
rather than a red parenthesis nobody exports cleanly.  Rates are written as
real numbers with a percent number-format -- never as the string "18%" and
never as the raw 0.18 sitting under a General format, which is how a 0.18%
rate ends up in a return.
"""

import os
import sys
import glob
from decimal import Decimal
from collections import OrderedDict, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import mtr_parser as M

ZERO = Decimal('0')
MONEY = '#,##0.00;-#,##0.00'
PCT = '0.00%'
INT = '#,##0'

HEAD_FILL = PatternFill('solid', fgColor='1F3864')
HEAD_FONT = Font(color='FFFFFF', bold=True, size=10)
SUB_FILL = PatternFill('solid', fgColor='D9E2F3')
BAD_FILL = PatternFill('solid', fgColor='FCE4E4')
WARN_FILL = PatternFill('solid', fgColor='FFF2CC')
THIN = Border(*[Side(style='thin', color='BFBFBF')] * 4)


def f(d):
    """Decimal -> float only at the very edge, because openpyxl cannot store a
    Decimal.  Every sum upstream of this line is exact."""
    return float(d) if isinstance(d, Decimal) else d


# --------------------------------------------------------------------------- #
# aggregation
# --------------------------------------------------------------------------- #

def blank():
    return {
        'rows': 0, 'sale_rows': 0, 'return_rows': 0, 'cancel_rows': 0,
        'gross_taxable': ZERO, 'return_taxable': ZERO,
        'gross_invoice': ZERO, 'return_invoice': ZERO,
        'cgst': ZERO, 'sgst': ZERO, 'utgst': ZERO, 'igst': ZERO, 'cess': ZERO,
        'total_tax': ZERO,
        'tcs_cgst': ZERO, 'tcs_sgst': ZERO, 'tcs_utgst': ZERO, 'tcs_igst': ZERO,
        'units': ZERO, 'promo_tax': ZERO,
        'einv_cancel_rows': 0, 'einv_cancel_net': ZERO,
    }


def accumulate(acc, r):
    acc['rows'] += 1
    kind = r['kind']
    if kind == 'SALE':
        acc['sale_rows'] += 1
        acc['gross_taxable'] += r['tax_exclusive_gross']
        acc['gross_invoice'] += r['invoice_amount']
    elif kind == 'RETURN':
        acc['return_rows'] += 1
        acc['return_taxable'] += r['tax_exclusive_gross']
        acc['return_invoice'] += r['invoice_amount']
    elif kind == 'CANCEL':
        acc['cancel_rows'] += 1
    elif kind in ('TRANSFER', 'TRANSFER_CANCEL'):
        # A stock transfer has no sale/return polarity; it is all "gross" so
        # that the transfer block foots to the file.
        acc['sale_rows'] += 1
        acc['gross_taxable'] += r['tax_exclusive_gross']
        acc['gross_invoice'] += r['invoice_amount']
    elif kind == 'EINVOICE_CANCEL':
        # A cancelled e-invoice is emitted as a +/- pair.  Counted, and its net
        # shown, so that "nets to nil" is visible rather than assumed.
        acc['einv_cancel_rows'] += 1
        acc['einv_cancel_net'] += r['invoice_amount']

    for h in ('cgst', 'sgst', 'utgst', 'igst', 'cess'):
        acc[h] += r['tax_by_head'][h]
    acc['total_tax'] += r['total_tax_amount']
    for h in ('cgst', 'sgst', 'utgst', 'igst'):
        acc['tcs_' + h] += r['tcs'][h]
    acc['units'] += r.get('quantity_signed', ZERO)
    acc['promo_tax'] += r.get('promo_tax_total', ZERO)
    return acc


def aggregate(docs):
    """(period_label, report_type) -> totals, in report-period order."""
    buckets = OrderedDict()
    for d in docs:
        y, mo, label = M.period_of(d['source_file'])
        sort_key = (y or 0, mo or 0)
        for r in d['rows']:
            k = (sort_key, label, r['report_type'])
            accumulate(buckets.setdefault(k, blank()), r)
    return OrderedDict(sorted(buckets.items(), key=lambda kv: (kv[0][0], kv[0][2])))


def aggregate_states(docs):
    """(period, origin, destination) -> totals + which heads were actually used."""
    out = OrderedDict()
    for d in docs:
        y, mo, label = M.period_of(d['source_file'])
        for r in d['rows']:
            if r['report_type'] == 'STOCK_TRANSFER':
                continue
            if r['total_tax_amount'] == 0 and r['tax_exclusive_gross'] == 0:
                continue       # a Cancel line carries no supply at all
            k = ((y or 0, mo or 0), label, r['origin_state'] or '(blank)',
                 r['destination_state'] or '(blank)')
            b = out.setdefault(k, blank())
            b['intra'] = r['is_intra_state']
            accumulate(b, r)
    return OrderedDict(sorted(out.items(), key=lambda kv: (kv[0][0], kv[0][2], kv[0][3])))


# --------------------------------------------------------------------------- #
# sheet helpers
# --------------------------------------------------------------------------- #

def header(ws, cols, row=1):
    for i, (title, _w, _fmt) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=title)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, THIN
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for i, (_t, w, _fmt) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    ws.auto_filter.ref = '%s%d:%s%d' % (get_column_letter(1), row,
                                        get_column_letter(len(cols)), row)


def put(ws, r, values, cols, fill=None, bold=False):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=r, column=i, value=f(v))
        fmt = cols[i - 1][2]
        if fmt:
            c.number_format = fmt
        c.border = THIN
        if fill:
            c.fill = fill
        if bold:
            c.font = Font(bold=True)
    return r + 1


# --------------------------------------------------------------------------- #
# sheets
# --------------------------------------------------------------------------- #

SUMMARY_COLS = [
    ('Month', 14, None), ('Type', 9, None),
    ('Rows', 8, INT), ('Sale rows', 10, INT), ('Return rows', 11, INT), ('Cancel rows', 11, INT),
    ('Gross taxable', 15, MONEY), ('Returns taxable', 15, MONEY), ('Net taxable', 15, MONEY),
    ('CGST', 12, MONEY), ('SGST', 12, MONEY), ('UTGST', 10, MONEY), ('IGST', 13, MONEY),
    ('Cess', 10, MONEY), ('Head sum', 13, MONEY), ('Promo tax', 12, MONEY),
    ('Total tax', 13, MONEY), ('Net invoice value', 16, MONEY),
    ('TCS CGST', 11, MONEY), ('TCS SGST', 11, MONEY), ('TCS IGST', 11, MONEY),
    ('TCS total', 12, MONEY),
    ('Net units', 10, INT),
    ('E-inv cancel rows', 12, INT), ('E-inv cancel net', 14, MONEY),
]


def sheet_summary(wb, docs, buckets, failures):
    ws = wb.create_sheet('Summary')
    ws.cell(row=1, column=1, value='Amazon MTR sales register -- SIMK LABELS PRIVATE LIMITED (Koparo)')\
        .font = Font(bold=True, size=13)
    ws.cell(row=2, column=1,
            value=('Returns arrive already negative from Amazon and are dated by credit-note date. '
                   'Net taxable = gross + returns. Stock transfers are excluded from turnover.'))\
        .font = Font(italic=True, size=9)
    note = ws.cell(row=3, column=1,
                   value=('%d verification failures -- see the Exceptions sheet.' % len(failures))
                   if failures else 'All verification checks passed.')
    note.font = Font(bold=True, size=10, color='9C0006' if failures else '006100')

    r = 5
    header(ws, SUMMARY_COLS, row=r)
    r += 1

    per_month = defaultdict(lambda: blank())
    sales_start = r
    for (sk, label, rtype), b in buckets.items():
        if rtype == 'STOCK_TRANSFER':
            continue
        net_taxable = b['gross_taxable'] + b['return_taxable']
        net_invoice = b['gross_invoice'] + b['return_invoice']
        tcs = b['tcs_cgst'] + b['tcs_sgst'] + b['tcs_utgst'] + b['tcs_igst']
        r = put(ws, r, [label, rtype, b['rows'], b['sale_rows'], b['return_rows'],
                        b['cancel_rows'], b['gross_taxable'], b['return_taxable'],
                        net_taxable, b['cgst'], b['sgst'], b['utgst'], b['igst'],
                        b['cess'],
                        b['cgst'] + b['sgst'] + b['utgst'] + b['igst'] + b['cess'],
                        b['promo_tax'], b['total_tax'], net_invoice,
                        b['tcs_cgst'], b['tcs_sgst'], b['tcs_igst'], tcs,
                        b['units'], b['einv_cancel_rows'], b['einv_cancel_net']], SUMMARY_COLS)
        m = per_month[(sk, label)]
        for k in blank():
            m[k] = m[k] + b[k]

    # Per-month B2B+B2C line, then the grand total.  Both are computed from the
    # same rows as the lines above, not re-read, so they cannot drift.
    r += 1
    ws.cell(row=r, column=1, value='Month totals (B2B + B2C)').font = Font(bold=True)
    r += 1
    header(ws, SUMMARY_COLS, row=r)
    r += 1
    grand = blank()
    for (sk, label), b in sorted(per_month.items()):
        net_taxable = b['gross_taxable'] + b['return_taxable']
        net_invoice = b['gross_invoice'] + b['return_invoice']
        tcs = b['tcs_cgst'] + b['tcs_sgst'] + b['tcs_utgst'] + b['tcs_igst']
        r = put(ws, r, [label, 'ALL', b['rows'], b['sale_rows'], b['return_rows'],
                        b['cancel_rows'], b['gross_taxable'], b['return_taxable'],
                        net_taxable, b['cgst'], b['sgst'], b['utgst'], b['igst'],
                        b['cess'],
                        b['cgst'] + b['sgst'] + b['utgst'] + b['igst'] + b['cess'],
                        b['promo_tax'], b['total_tax'], net_invoice,
                        b['tcs_cgst'], b['tcs_sgst'], b['tcs_igst'], tcs,
                        b['units'], b['einv_cancel_rows'], b['einv_cancel_net']],
                SUMMARY_COLS, fill=SUB_FILL)
        for k in blank():
            grand[k] = grand[k] + b[k]
    tcs = grand['tcs_cgst'] + grand['tcs_sgst'] + grand['tcs_utgst'] + grand['tcs_igst']
    r = put(ws, r, ['Apr-Aug 2026', 'TOTAL', grand['rows'], grand['sale_rows'],
                    grand['return_rows'], grand['cancel_rows'], grand['gross_taxable'],
                    grand['return_taxable'], grand['gross_taxable'] + grand['return_taxable'],
                    grand['cgst'], grand['sgst'], grand['utgst'], grand['igst'],
                    grand['cess'],
                    grand['cgst'] + grand['sgst'] + grand['utgst'] + grand['igst'] + grand['cess'],
                    grand['promo_tax'], grand['total_tax'],
                    grand['gross_invoice'] + grand['return_invoice'],
                    grand['tcs_cgst'], grand['tcs_sgst'], grand['tcs_igst'], tcs,
                    grand['units'], grand['einv_cancel_rows'], grand['einv_cancel_net']],
            SUMMARY_COLS, fill=SUB_FILL, bold=True)

    # Stock transfers, deliberately below and apart.
    r += 2
    ws.cell(row=r, column=1,
            value='Stock transfers (own GSTIN to own GSTIN -- NOT turnover)').font = Font(bold=True)
    r += 1
    st_cols = [('Month', 14, None), ('Rows', 8, INT), ('Taxable value', 15, MONEY),
               ('CGST', 12, MONEY), ('SGST', 12, MONEY), ('IGST', 13, MONEY),
               ('Total tax', 13, MONEY), ('Invoice value', 15, MONEY), ('Units', 10, INT)]
    for i, (t, w, _fmt) in enumerate(st_cols, start=1):
        c = ws.cell(row=r, column=i, value=t)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, THIN
    r += 1
    for (sk, label, rtype), b in buckets.items():
        if rtype != 'STOCK_TRANSFER':
            continue
        r = put(ws, r, [label, b['rows'], b['gross_taxable'] + b['return_taxable'],
                        b['cgst'], b['sgst'], b['igst'], b['total_tax'],
                        b['gross_invoice'] + b['return_invoice'], b['units']], st_cols)
    return ws


STATE_COLS = [
    ('Month', 14, None), ('Ship-from state', 20, None), ('Ship-to / bill-to state', 22, None),
    ('Supply', 11, None), ('Rows', 8, INT),
    ('Gross taxable', 15, MONEY), ('Returns taxable', 15, MONEY), ('Net taxable', 15, MONEY),
    ('CGST', 12, MONEY), ('SGST', 12, MONEY), ('UTGST', 10, MONEY), ('IGST', 13, MONEY),
    ('Cess', 10, MONEY), ('Total tax', 13, MONEY),
    ('Effective tax rate', 14, PCT),
    ('Heads used', 16, None), ('Head matches supply', 18, None),
]


def sheet_states(wb, state_buckets):
    """Ship-from vs ship-to is the whole point of this sheet.

    Destination is the BILL-TO state when the buyer is registered and gave one,
    because a bill-to/ship-to supply is taxed where the buyer is, not where the
    carton went.  Nine B2B rows in these files are taxed that way and would look
    like errors on a naive ship-to grouping.
    """
    ws = wb.create_sheet('State-wise')
    ws.cell(row=1, column=1, value='Monthly state-wise supply').font = Font(bold=True, size=12)
    ws.cell(row=2, column=1,
            value=('Destination = Bill To state where the buyer gave a GSTIN (bill-to/ship-to), '
                   'else Ship To state. Same state => CGST+SGST/UTGST; different => IGST.'))\
        .font = Font(italic=True, size=9)
    r = 4
    header(ws, STATE_COLS, row=r)
    r += 1
    for (sk, label, origin, dest), b in state_buckets.items():
        intra = b['intra']
        used = [h.upper() for h in ('cgst', 'sgst', 'utgst', 'igst', 'cess') if b[h] != 0]
        expected_ok = (('IGST' not in used) if intra
                       else not ({'CGST', 'SGST', 'UTGST'} & set(used)))
        net_taxable = b['gross_taxable'] + b['return_taxable']
        rate = (b['total_tax'] / net_taxable) if net_taxable else ZERO
        r = put(ws, r, [label, origin, dest, 'Intra-state' if intra else 'Inter-state',
                        b['rows'], b['gross_taxable'], b['return_taxable'], net_taxable,
                        b['cgst'], b['sgst'], b['utgst'], b['igst'], b['cess'],
                        b['total_tax'], rate, '+'.join(used) or '(none)',
                        'OK' if expected_ok else 'MISMATCH'],
                STATE_COLS, fill=None if expected_ok else BAD_FILL)
    return ws


DETAIL_COLS = [
    ('Source file', 30, None), ('Line', 7, INT), ('Type', 7, None),
    ('Transaction type', 15, None), ('Kind', 15, None),
    ('Doc date', 12, None), ('Invoice date', 12, None), ('Credit note date', 14, None),
    ('Invoice no', 14, None), ('Credit note no', 14, None), ('Order id', 20, None),
    ('Seller GSTIN', 17, None), ('Customer GSTIN', 17, None), ('Buyer name', 22, None),
    ('SKU', 22, None), ('ASIN', 12, None), ('HSN', 10, None), ('Item description', 34, None),
    ('Qty', 6, INT), ('Qty signed', 9, INT),
    ('Ship-from state', 17, None), ('Ship-to state', 17, None), ('Bill-to state', 17, None),
    ('Supply', 11, None), ('Warehouse', 10, None), ('Channel', 8, None),
    ('Taxable value', 14, MONEY), ('CGST', 11, MONEY), ('SGST', 11, MONEY),
    ('UTGST', 9, MONEY), ('IGST', 12, MONEY), ('Cess', 9, MONEY),
    ('Total tax', 12, MONEY), ('Invoice amount', 14, MONEY),
    ('CGST rate', 9, PCT), ('SGST rate', 9, PCT), ('IGST rate', 9, PCT),
    ('TCS total', 10, MONEY), ('IRN', 40, None), ('IRN status', 12, None),
]


def sheet_detail(wb, docs):
    ws = wb.create_sheet('Detail')
    r = 1
    header(ws, DETAIL_COLS, row=r)
    r += 1
    for d in docs:
        if d['report_type'] == 'STOCK_TRANSFER':
            continue
        for x in d['rows']:
            dt = lambda k: x[k].strftime('%Y-%m-%d') if x.get(k) else ''
            r = put(ws, r, [
                x['source_file'], x['line_no'], x['report_type'], x['transaction_type'],
                x['kind'], dt('doc_date'), dt('invoice_date'), dt('credit_note_date'),
                x['invoice_number'], x['credit_note_no'], x['order_id'],
                x['seller_gstin'], x['customer_gstin'], x['buyer_name'],
                x['sku'], x['asin'], x['hsn'], x['item_description'][:120],
                x['quantity'], x['quantity_signed'],
                x['ship_from_state'], x['ship_to_state'], x['bill_to_state'],
                'Intra-state' if x['is_intra_state'] else 'Inter-state',
                x['warehouse_id'], x['fulfillment_channel'],
                x['tax_exclusive_gross'],
                x['tax_by_head']['cgst'], x['tax_by_head']['sgst'], x['tax_by_head']['utgst'],
                x['tax_by_head']['igst'], x['tax_by_head']['cess'],
                x['total_tax_amount'], x['invoice_amount'],
                x['rates']['cgst'], x['rates']['sgst'], x['rates']['igst'],
                x['tcs_total'], x['irn_number'], x['irn_status'],
            ], DETAIL_COLS)
    return ws


TRANSFER_COLS = [
    ('Source file', 30, None), ('Line', 7, INT), ('Transaction type', 18, None),
    ('Invoice date', 12, None), ('Invoice no', 14, None), ('Order id', 20, None),
    ('Supplier GSTIN', 17, None), ('Receiver GSTIN', 17, None),
    ('Ship-from FC', 12, None), ('Ship-from state', 17, None),
    ('Ship-to FC', 12, None), ('Ship-to state', 17, None), ('Supply', 11, None),
    ('SKU', 22, None), ('HSN', 10, None), ('Qty', 6, INT),
    ('Taxable value', 14, MONEY), ('CGST', 11, MONEY), ('SGST', 11, MONEY),
    ('IGST', 12, MONEY), ('Cess', 9, MONEY), ('Total tax', 12, MONEY),
    ('Invoice value', 14, MONEY), ('IGST rate', 9, PCT),
]


def sheet_transfers(wb, docs):
    ws = wb.create_sheet('Stock Transfers')
    ws.cell(row=1, column=1,
            value=('FC-to-FC movements between Koparo\'s own registrations. '
                   'Not sales -- excluded from every turnover figure in this book.'))\
        .font = Font(italic=True, size=9)
    r = 3
    header(ws, TRANSFER_COLS, row=r)
    r += 1
    for d in docs:
        if d['report_type'] != 'STOCK_TRANSFER':
            continue
        for x in d['rows']:
            r = put(ws, r, [
                x['source_file'], x['line_no'], x['transaction_type'],
                x['invoice_date'].strftime('%Y-%m-%d') if x['invoice_date'] else '',
                x['invoice_number'], x['order_id'],
                x['seller_gstin'], x['receiver_gstin'],
                x['ship_from_fc'], x['ship_from_state'],
                x['ship_to_fc'], x['ship_to_state'],
                'Intra-state' if x['is_intra_state'] else 'Inter-state',
                x['sku'], x['hsn'], x['quantity'],
                x['tax_exclusive_gross'], x['tax_by_head']['cgst'], x['tax_by_head']['sgst'],
                x['tax_by_head']['igst'], x['tax_by_head']['cess'], x['total_tax_amount'],
                x['invoice_amount'], x['rates']['igst'],
            ], TRANSFER_COLS)
    return ws


EXC_COLS = [
    ('Source file', 32, None), ('Line', 7, INT), ('Check', 60, None),
    ('Invoice no', 14, None), ('Order id', 20, None), ('Transaction type', 16, None),
    ('Amounts', 90, None),
]


def sheet_exceptions(wb, failures):
    ws = wb.create_sheet('Exceptions')
    ws.cell(row=1, column=1,
            value='Every row that failed a check. Nothing here was dropped from the other sheets.')\
        .font = Font(bold=True, size=10, color='9C0006')
    r = 3
    header(ws, EXC_COLS, row=r)
    r += 1
    for x in sorted(failures, key=lambda z: (z['file'], z['line_no'], z['check'])):
        r = put(ws, r, [x['file'], x['line_no'], x['check'], x['invoice_number'],
                        x['order_id'], x['transaction_type'],
                        '; '.join('%s=%s' % kv for kv in x['amounts'].items())],
                EXC_COLS, fill=BAD_FILL)
    if not failures:
        ws.cell(row=r, column=1, value='No failures.').fill = WARN_FILL
    return ws


# --------------------------------------------------------------------------- #

def build(folder, out_path):
    paths = sorted(glob.glob(os.path.join(folder, '*.csv')))
    if not paths:
        raise SystemExit('no CSVs under %s' % folder)
    docs = [M.parse(p) for p in paths]
    failures = [x for d in docs for x in d['failures']]

    wb = Workbook()
    wb.remove(wb.active)
    sheet_summary(wb, docs, aggregate(docs), failures)
    sheet_states(wb, aggregate_states(docs))
    sheet_detail(wb, docs)
    sheet_transfers(wb, docs)
    sheet_exceptions(wb, failures)
    wb.save(out_path)
    return docs, failures


if __name__ == '__main__':
    folder = sys.argv[1] if len(sys.argv) > 1 else '.'
    out = sys.argv[2] if len(sys.argv) > 2 else 'amazon_mtr_sales_register.xlsx'
    docs, failures = build(folder, out)
    print('wrote %s' % out)
    print('%d files, %d rows, %d verification failures'
          % (len(docs), sum(d['row_count'] for d in docs), len(failures)))
