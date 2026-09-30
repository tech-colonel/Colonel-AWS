import openpyxl

GEN = "/Users/apple/Documents/Colonel-AWS/new-backend/scripts/dchica-shopify/output_generated.xlsx"
EXP = "/Users/apple/Documents/Colonel-AWS/test files receivables/shopify dchica/Total sales breakdown (P) With Order Name Jul-26 Working.xlsx"

wb_gen = openpyxl.load_workbook(GEN, read_only=True, data_only=True)
print("Generated sheets:", wb_gen.sheetnames)
ws_gen = wb_gen['Working']
print(wb_gen.sheetnames)

gen_header = [c.value for c in next(ws_gen.iter_rows(min_row=1, max_row=1))]
print("gen header:", gen_header)

gen_rows = list(ws_gen.iter_rows(min_row=2, values_only=True))
print("gen rows:", len(gen_rows))

wb_exp = openpyxl.load_workbook(EXP, read_only=True, data_only=True)
ws_exp = wb_exp['Total sales breakdown (P) With ']
exp_header = [c.value for c in next(ws_exp.iter_rows(min_row=4, max_row=4))]
exp_rows = list(ws_exp.iter_rows(min_row=5, values_only=True))
print("exp rows:", len(exp_rows))

# column index maps
def idx(header, name):
    return header.index(name)

gi = {n: idx(gen_header, n) for n in ['Order ID', 'Product variant SKU', 'Product Type', 'HSN Code', 'GST Rate', 'Net Qty', 'Invoice Value', 'Taxable', 'IGST', 'CGST', 'SGST']}
ei = {
    'Order ID': idx(exp_header, 'Order ID'),
    'Product variant SKU': idx(exp_header, 'Product variant SKU'),
    'Product Type': idx(exp_header, 'Product type'),
    'HSN Code': idx(exp_header, 'HSN Code'),
    'GST Rate': idx(exp_header, 'GST Rate'),
    'Net Qty': idx(exp_header, 'Net Qty'),
    'Invoice Value': idx(exp_header, 'Invoice Value'),
    'Taxable': idx(exp_header, 'Taxable'),
    'IGST': idx(exp_header, 'IGST'),
    'CGST': idx(exp_header, 'CGST'),
    'SGST': idx(exp_header, 'SGST'),
}

assert len(gen_rows) == len(exp_rows), f"row count mismatch {len(gen_rows)} vs {len(exp_rows)}"

def close(a, b, tol=0.5):
    try:
        return abs(float(a or 0) - float(b or 0)) <= tol
    except (TypeError, ValueError):
        return str(a or '').strip() == str(b or '').strip()

cols = ['Product Type', 'HSN Code', 'GST Rate', 'Net Qty', 'Invoice Value', 'Taxable', 'IGST', 'CGST', 'SGST']
mismatch_counts = {c: 0 for c in cols}
examples = {c: [] for c in cols}

for r_idx, (g, e) in enumerate(zip(gen_rows, exp_rows)):
    for c in cols:
        gv = g[gi[c]]
        ev = e[ei[c]]
        if c == 'Product Type':
            ok = str(gv or '').strip().lower() == str(ev or '').strip().lower()
        else:
            ok = close(gv, ev)
        if not ok:
            mismatch_counts[c] += 1
            if len(examples[c]) < 5:
                examples[c].append((r_idx + 5, g[gi['Order ID']], g[gi['Product variant SKU']], gv, ev))

n = len(gen_rows)
print("\n=== VALIDATION REPORT ===")
for c in cols:
    mc = mismatch_counts[c]
    print(f"{c}: {n - mc}/{n} match ({100*(n-mc)/n:.2f}%)")
    for ex in examples[c]:
        print("   row", ex[0], "order", ex[1], "sku", ex[2], "generated=", ex[3], "expected=", ex[4])
