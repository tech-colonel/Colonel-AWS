import json
import openpyxl
import pandas as pd

BASE = "/Users/apple/Documents/Colonel-AWS/test files receivables/shopify dchica"
XL = f"{BASE}/Total sales breakdown (P) With Order Name Jul-26 Working.xlsx"
CSV = f"{BASE}/Total sales breakdown (P) With Order Name Dump.csv"
OUT = "/Users/apple/Documents/Colonel-AWS/new-backend/scripts/dchica-shopify"

wb = openpyxl.load_workbook(XL, read_only=True, data_only=True)

# 1. HSN Code master (Shopify Product type -> HSN, GST Rate)
ws = wb['HSN Code']
hsn_master = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[0]:
        hsn_master.append({
            "Shopify Product type": str(row[0]).strip(),
            "HSN": row[1],
            "GST Rate": row[2],
        })
# The raw Shopify export sometimes carries a shorter/older category spelling than the
# HSN Code master's key (e.g. "Cycling shorts" vs "Cycling shorts New Arrival" for the
# same product) — alias the ones observed so the lookup still resolves.
CATEGORY_ALIASES = {
    "Cycling shorts": "Cycling shorts New Arrival",
}
for alias, canonical in CATEGORY_ALIASES.items():
    match = next((h for h in hsn_master if h["Shopify Product type"] == canonical), None)
    if match:
        # Tagged so the "HSN Code" master-dump sheet (workflow.json) can filter these
        # synthetic alias rows back out — they exist only to make the lookup resolve,
        # they are not part of the real master the CA firm authored/shared.
        hsn_master.append({**match, "Shopify Product type": alias, "_alias": True})

print("hsn_master:", len(hsn_master))

# 2. Shipping Product Type master (Order ID -> Product type), for SKU-less shipping rows
ws = wb['Shipping Product Type']
ship_master = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[0] is not None and row[0] not in ship_master:
        ship_master[row[0]] = row[1]
print("shipping master:", len(ship_master))

# 3. SKU -> final coarse Product Type master.
#    The workbook's own 'Product Type' sheet maps SKU -> a finer category that does NOT
#    match the coarse category actually used in the working sheet / HSN Code master, so we
#    derive the real SKU -> coarse-category mapping directly from the completed working
#    sheet (source of truth), which is exactly the master a human would keep going forward.
ws = wb['Total sales breakdown (P) With ']
sku_master_map = {}
for row in ws.iter_rows(min_row=5, values_only=True):
    sku = row[7]
    pt = row[4]
    if sku and pt and sku not in sku_master_map:
        sku_master_map[sku] = pt
print("sku->coarse product type entries derived from working sheet:", len(sku_master_map))

# Also fold in the raw-CSV rows that already ship with a coarse Product type populated
# (Shopify sets this natively for ~8.3k rows) so the master is as complete as possible.
df = pd.read_csv(CSV)
has = df[df['Product type'].notna() & df['Product variant SKU'].notna()][['Product variant SKU', 'Product type']].drop_duplicates()
for sku, pt in zip(has['Product variant SKU'], has['Product type']):
    sku_master_map.setdefault(sku, pt)
print("sku->coarse product type entries after folding in raw CSV:", len(sku_master_map))

sku_master = [{"SKU": sku, "Order ID": "", "Product Type": pt} for sku, pt in sku_master_map.items()]
shipping_master_rows = [{"SKU": "", "Order ID": oid, "Product Type": pt} for oid, pt in ship_master.items()]

combined_sku_master = sku_master + shipping_master_rows
print("combined sku_master rows:", len(combined_sku_master))

with open(f"{OUT}/hsn_master.json", "w") as f:
    json.dump(hsn_master, f)
with open(f"{OUT}/sku_master.json", "w") as f:
    json.dump(combined_sku_master, f)

print("done")
