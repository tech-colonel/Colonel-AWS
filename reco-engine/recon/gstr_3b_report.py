"""GSTR-1/3B vs Books — the reconciliation report (one tab) + the input sheets.

Layout approved by the accountant (blueprint → template):

    GTS                       ┐
    (−) Returns               │  one table each: months down (Apr … Mar + Total),
    (−) Credit Notes  *       │  registrations across (Taxable Value, IGST, CGST,
    Net Sales                 │  SGST per registration) + Total — all registrations
    GSTR-1 / 3B               │
    Difference                │  Difference     = GSTR-1/3B − Net Sales
    (−) Inter Sales   *       │  Net Difference = Difference − Inter Sales
    Net Difference    *       ┘
    * only when that data exists — without it the report is the same, minus the table.

Every input figure carries a note (where it comes from, rows, top ledgers, file) and a
link to exactly those rows on the input tab; every total shows the range it combines.
Net Sales / Difference / Net Difference and all totals are live Excel formulas.
"""
from __future__ import annotations

import re
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.hyperlink import Hyperlink

from .gstr_1_vs_books import _FY_MONTHS

HEADS = (("net", "Taxable Value"), ("igst", "IGST"), ("cgst", "CGST"), ("sgst", "SGST"))
NUM = '_ * #,##0_ ;_ * \\-#,##0_ ;_ * "-"??_ ;_ @_ '
GREY, BAND = "7F7F7F", "F5F7FA"
SECTION = {                                         # title bar, state header, total tint
    "GTS":            ("1F4E78", "2E75B6", "D9E1F2"),
    "Returns":        ("843C0C", "C55A11", "FBE5D6"),
    "Credit Notes":   ("4B1E6B", "7030A0", "E4DFEC"),
    "Net Sales":      ("375623", "548235", "E2EFDA"),
    "Return":         ("0B5563", "138D90", "D6EEEE"),
    "Difference":     ("7F6000", "BF8F00", "FFF2CC"),
    "Inter Sales":    ("3B3838", "767171", "EDEDED"),
    "Net Difference": ("7B0000", "C00000", "FADADD"),
}
NOTE_NAME = {"GTS": "GTS", "Returns": "Returns", "Credit Notes": "Credit Notes", "Net Sales": "Net Sales",
             "Return": "GSTR-1 / 3B", "Difference": "Difference", "Inter Sales": "Inter Sales",
             "Net Difference": "Net Difference"}
TITLES = {
    "GTS": "GTS — GROSS TOTAL SALES (BOOKS)",
    "Returns": "(−) RETURNS",
    "Credit Notes": "(−) CREDIT NOTES",
    "Net Sales": "NET SALES  (GTS − Returns{cn})",
    "Return": "GSTR-1 / 3B SUMMARY",
    "Difference": "DIFFERENCE  (GSTR-1/3B − Net Sales)",
    "Inter Sales": "(−) INTER SALES",
    "Net Difference": "NET DIFFERENCE  (Difference − Inter Sales)",
}
_thin = Side(style="thin", color="BFBFBF")
_sep = Side(style="medium", color="404040")
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _font(**k):
    return Font(name="Calibri", size=k.pop("size", 11), **k)


def _fill(c):
    return PatternFill("solid", fgColor=c)


def _is_credit_note(r) -> bool:
    if r["_block"]["kind"] == "credit_note":
        return True
    return bool(re.search(r"credit\s*note", f"{r['source']} {r['voucher_type']}", re.I))


def _category(r) -> str:
    if r["category"] == "interbranch":
        return "Inter Sales"
    if r["category"] == "returns":
        return "Credit Notes" if _is_credit_note(r) else "Returns"
    return "GTS"


def build_report(res: dict) -> Workbook:
    summ, books = res["summ"], res["books"]
    regs = list(res["states"])                               # registrations, report order
    labels = summ["labels"]
    fy_start = res["fy_start"]
    months = [date(fy_start if datetime_month(m) >= 4 else fy_start + 1, datetime_month(m), 1) for m in _FY_MONTHS]

    # ---- input tab first, sorted so each report cell's rows are one contiguous block
    rows = sorted(books["rows"], key=lambda r: (regs.index(r["state"]) if r["state"] in regs else 99,
                                                ["GTS", "Returns", "Credit Notes", "Inter Sales"].index(_category(r)),
                                                _FY_MONTHS.index(r["month"]) if r["month"] in _FY_MONTHS else 99,
                                                r["row"]))
    cell_rows: dict[tuple, list] = defaultdict(list)        # (cat, reg, month) -> [input-tab row numbers]
    agg: dict[tuple, dict] = defaultdict(lambda: {k: 0.0 for k, _ in HEADS})
    ledgers: dict[tuple, dict] = defaultdict(lambda: defaultdict(float))
    src_rows: dict[tuple, list] = defaultdict(list)          # (cat, reg, month) -> original register rows
    for i, r in enumerate(rows):
        key = (_category(r), r["state"], r["month"])
        cell_rows[key].append(i + 2)
        src_rows[key].append(f"{r['sheet']}!{r['row']}")
        for k, _ in HEADS:
            agg[key][k] += r["amt"][k]
        for h, v in r["ledgers"].items():
            ledgers[key][h] += abs(v)
    have_cn = any(abs(v) > 0.005 for (c, _s, _m), d in agg.items() if c == "Credit Notes" for v in d.values())
    have_ib = any(abs(v) > 0.005 for (c, _s, _m), d in agg.items() if c == "Inter Sales" for v in d.values())

    wb = Workbook()
    ws = wb.active
    ws.title = "1-3B vs Books"
    ws.sheet_view.showGridLines = False
    data_ws = wb.create_sheet("Sales Register")
    ret_ws = wb.create_sheet("GSTR-1 & 3B")
    _write_inputs(data_ws, ret_ws, rows, res, labels)
    ret_row = {(d["Registration"], d["Month"]): i + 2 for i, d in enumerate(res["detail"])}
    ret_file = {(d["Registration"], d["Month"]): d for d in res["detail"]}

    groups = regs + ["Total"]
    ncols = 1 + 4 * len(groups)

    def edge(col):
        first = col >= 2 and (col - 2) % 4 == 0
        return Border(left=_sep if first else _thin, right=_sep if col == ncols else _thin, top=_thin, bottom=_thin)

    who = res["company"] or "Sales Register"
    ws["A1"] = "GSTR-1 / 3B vs Books — Reconciliation"
    ws["A1"].font = _font(size=14, bold=True)
    ws["A2"] = (f"{who}    |    FY {res['fy']}    |    Registrations: "
                + ", ".join(f"{labels.get(s, s)} ({res['gstin_of'].get(s, '—')})" for s in regs))
    ws["A2"].font = _font(size=9, color="595959")
    ws["A3"] = ("Hover a figure (or right-click → Show Note) to see where it comes from; click a Books or return figure "
                "to jump to its rows. A Total shows the range it combines. Return = GSTR-3B 3.1(a)+(b); where no 3B was "
                "uploaded, that month's GSTR-1.")
    ws["A3"].font = _font(size=9, italic=True, color="595959")

    order = ["GTS", "Returns"] + (["Credit Notes"] if have_cn else []) + ["Net Sales", "Return", "Difference"] \
        + (["Inter Sales", "Net Difference"] if have_ib else [])
    first_row, r = {}, 5
    for key in order:                                        # every table: bar + 2 header rows + 13 rows + 2 gap
        first_row[key] = r + 3
        r += 18

    r = 5
    for key in order:
        dark, mid, tint = SECTION[key]
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        t = ws.cell(row=r, column=1, value="  " + TITLES[key].format(cn=" − Credit Notes" if have_cn else ""))
        t.font, t.fill = _font(bold=True, size=12, color="FFFFFF"), _fill(dark)
        t.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[r].height = 20
        r += 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r + 1, end_column=1)
        h = ws.cell(row=r, column=1, value="Month")
        h.font, h.fill, h.alignment = _font(bold=True, color="FFFFFF"), _fill(mid), _CENTER
        ws.cell(row=r + 1, column=1).border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
        for g, st in enumerate(groups):
            c0 = 2 + 4 * g
            ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 3)
            name = "Total — all registrations" if st == "Total" else \
                f"{labels.get(st, st)}" + (f" — {res['gstin_of'][st]}" if res["gstin_of"].get(st) else "")
            x = ws.cell(row=r, column=c0, value=name)
            x.font, x.fill, x.alignment = _font(bold=True, color="FFFFFF"), _fill(dark if st == "Total" else mid), _CENTER
            for j in range(4):
                ws.cell(row=r, column=c0 + j).border = edge(c0 + j)
            for j, (_k, hd) in enumerate(HEADS):
                y = ws.cell(row=r + 1, column=c0 + j, value=hd)
                y.font, y.fill, y.alignment, y.border = _font(size=10, bold=True, color="FFFFFF"), _fill(GREY), _CENTER, edge(c0 + j)
        r += 2
        top = r
        for mi, d in enumerate(months + [None]):
            is_total = d is None
            month = _FY_MONTHS[mi] if d else None
            m = ws.cell(row=r, column=1, value=d if d else "Total")
            m.number_format = "mmm-yy"
            m.font, m.alignment = _font(bold=True), _CENTER
            m.border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
            if is_total:
                m.fill = _fill(tint)
            when = d.strftime("%B %Y") if d else "April–March (FY total)"
            for g, st in enumerate(groups):
                for j, (k, hd) in enumerate(HEADS):
                    col = 2 + 4 * g + j
                    cl = L(col)
                    cell = ws.cell(row=r, column=col)
                    off = r - top
                    note, link = "", None
                    if is_total:
                        cell.value = f"=SUM({cl}{top}:{cl}{r - 1})"
                        note = f"RANGE: April–March (12 months)" + (", all registrations" if st == "Total" else f", {labels.get(st, st)}")
                    elif st == "Total":
                        cell.value = "=" + "+".join(f"{L(2 + 4 * s + j)}{r}" for s in range(len(regs)))
                        note = "RANGE: " + " + ".join(labels.get(s, s) for s in regs) + f", {when}"
                    elif key == "Net Sales":
                        cell.value = (f"={cl}{first_row['GTS'] + off}-{cl}{first_row['Returns'] + off}"
                                      + (f"-{cl}{first_row['Credit Notes'] + off}" if have_cn else ""))
                        note = "= GTS − Returns" + (" − Credit Notes" if have_cn else "") + " (same month, registration, tax head)"
                    elif key == "Difference":
                        cell.value = f"={cl}{first_row['Return'] + off}-{cl}{first_row['Net Sales'] + off}"
                        note = "= GSTR-1/3B − Net Sales"
                    elif key == "Net Difference":
                        cell.value = f"={cl}{first_row['Difference'] + off}-{cl}{first_row['Inter Sales'] + off}"
                        note = ("= Difference − Inter Sales. Inter sales are booked as sales in the return but are only "
                                "transfers between own registrations — what is left is the real gap in Books or GSTR-1/3B.")
                    elif key == "Return":
                        c = res["comp"][st][month]
                        g3 = c["g3b"]
                        cell.value = round(g3[k], 2) if g3 is not None else None
                        det = ret_file.get((labels.get(st, st), month))
                        if det:
                            note = (f"{det['Compared with']} ({det['Format']}) — {det.get('File') or 'OCTA sheet'}"
                                    + (" · Nil return" if det.get("Nil return") else "")
                                    + (" · GSTR-3B table 3.1(a)+(b)" if det["Compared with"] == "GSTR-3B"
                                       else " · GSTR-1 Total Liability (no GSTR-3B uploaded for this month)"))
                            rr_ = ret_row[(labels.get(st, st), month)]
                            link = f"'GSTR-1 & 3B'!A{rr_}:{L(len(res['detail'][0]))}{rr_}"
                        else:
                            note = "No GSTR-3B or GSTR-1 uploaded for this registration and month."
                    else:
                        kk = (key, st, month)
                        v = agg.get(kk, {}).get(k)
                        cell.value = round(v, 2) if v is not None and kk in agg else None
                        rr = cell_rows.get(kk, [])
                        if rr:
                            top_l = sorted(ledgers[kk].items(), key=lambda x: -x[1])[:3]
                            note = (f"{len(rr)} row(s) — Sales Register tab rows {rr[0]}–{rr[-1]}"
                                    f" (register: {', '.join(src_rows[kk][:4])}{' …' if len(src_rows[kk]) > 4 else ''})")
                            if k == "net" and top_l:
                                note += "\nTop ledgers: " + "; ".join(f"{h} {v2:,.0f}" for h, v2 in top_l)
                            link = f"'Sales Register'!A{rr[0]}:T{rr[-1]}"
                        else:
                            note = "No rows in the Books for this registration and month."
                    cell.number_format = NUM
                    cell.border = edge(col)
                    cell.font = _font(bold=is_total or st == "Total")
                    if is_total or st == "Total":
                        cell.fill = _fill(tint)
                    elif g % 2 == 1:
                        cell.fill = _fill(BAND)
                    head = f"{NOTE_NAME[key]} · {hd} · {'All registrations' if st == 'Total' else labels.get(st, st)} · {when}"
                    cell.comment = Comment(head + "\n" + note, "Colonel", width=340, height=120)
                    if link:                                  # in-workbook jump to the source rows
                        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=link,
                                                   tooltip="Go to the rows behind this figure")
                        cell.font = _font(bold=False, color="1F3864")
            r += 1
        r += 2

    ws.column_dimensions["A"].width = 12
    for c in range(2, ncols + 1):
        ws.column_dimensions[L(c)].width = 15
    ws.freeze_panes = "B5"
    return wb


def datetime_month(name: str) -> int:
    return ["January", "February", "March", "April", "May", "June", "July", "August", "September",
            "October", "November", "December"].index(name) + 1


def _write_inputs(data_ws, ret_ws, rows, res, labels):
    head_font, head_fill = _font(bold=True, color="FFFFFF"), _fill("1F4E78")
    cols = ["Registration", "Category", "Month", "Date", "Particulars", "Voucher Type", "Narration", "Nature",
            "Source", "Register sheet", "Register row", "State field", "How state was assigned",
            "Taxable Value", "IGST", "CGST", "SGST", "Total GST", "Gross", "Register Gross Total"]
    for i, h in enumerate(cols, 1):
        c = data_ws.cell(row=1, column=i, value=h)
        c.font, c.fill, c.alignment = head_font, head_fill, _CENTER
    for n, r in enumerate(rows, 2):
        a = r["amt"]
        gst = round(a["cgst"] + a["sgst"] + a["igst"], 2)
        vals = [labels.get(r["state"], r["state"]), _category(r), r["month"],
                r["date"].date() if r["date"] else None, r["particulars"], r["voucher_type"], r["narration"],
                r["nature"], r["source"], r["sheet"], r["row"], r["state_field"], r["state_why"],
                a["net"], a["igst"], a["cgst"], a["sgst"], gst, round(a["net"] + gst, 2), r["gross"]]
        for i, v in enumerate(vals, 1):
            c = data_ws.cell(row=n, column=i, value=v)
            if i == 4 and v:
                c.number_format = "dd-mmm-yy"
            elif isinstance(v, float):
                c.number_format = NUM
    widths = [16, 13, 11, 11, 34, 14, 40, 9, 18, 14, 10, 13, 34, 15, 13, 13, 13, 13, 15, 15]
    for i, w in enumerate(widths, 1):
        data_ws.column_dimensions[L(i)].width = w
    data_ws.freeze_panes = "A2"
    data_ws.auto_filter.ref = f"A1:{L(len(cols))}{max(2, len(rows) + 1)}"

    det = res["detail"]
    if det:
        hdr = list(det[0].keys())
        for i, h in enumerate(hdr, 1):
            c = ret_ws.cell(row=1, column=i, value=h)
            c.font, c.fill, c.alignment = head_font, _fill("0B5563"), _CENTER
        for n, d in enumerate(det, 2):
            for i, h in enumerate(hdr, 1):
                v = d[h]
                c = ret_ws.cell(row=n, column=i, value=v)
                if isinstance(v, float):
                    c.number_format = NUM
        for i in range(1, len(hdr) + 1):
            ret_ws.column_dimensions[L(i)].width = 16
        ret_ws.column_dimensions[L(hdr.index("File") + 1)].width = 40
        ret_ws.freeze_panes = "A2"
    else:
        ret_ws.cell(row=1, column=1, value="No GSTR-1 or GSTR-3B uploaded.")


def report_values(res: dict) -> dict:
    """The report's eight lines as numbers, for the web page — same categorisation and
    the same arithmetic as the Excel formulas (Net Sales = GTS − Returns − Credit Notes,
    Difference = GSTR-1/3B − Net Sales, Net Difference = Difference − Inter Sales)."""
    summ, books = res["summ"], res["books"]
    regs = list(res["states"])
    labels = summ["labels"]
    keys = [k for k, _ in HEADS]
    z = lambda: {k: 0.0 for k in keys}                                   # noqa: E731
    base = {c: {s: {m: z() for m in _FY_MONTHS} for s in regs} for c in ("GTS", "Returns", "Credit Notes", "Inter Sales")}
    for r in books["rows"]:
        if r["state"] in regs and r["month"] in _FY_MONTHS:
            for k in keys:
                base[_category(r)][r["state"]][r["month"]][k] += r["amt"][k]
    have_cn = any(abs(v) > 0.005 for s in regs for m in _FY_MONTHS for v in base["Credit Notes"][s][m].values())
    have_ib = any(abs(v) > 0.005 for s in regs for m in _FY_MONTHS for v in base["Inter Sales"][s][m].values())
    ret = {s: {m: ({k: (res["comp"][s][m]["g3b"] or {}).get(k, 0.0) for k in keys}
                   if res["comp"][s][m]["g3b"] is not None else None) for m in _FY_MONTHS} for s in regs}
    lines = ["GTS", "Returns"] + (["Credit Notes"] if have_cn else []) + ["Net Sales", "Return", "Difference"] \
        + (["Inter Sales", "Net Difference"] if have_ib else [])
    vals: dict = {ln: {} for ln in lines}
    for s in regs:
        for ln in lines:
            vals[ln][s] = {}
        for m in _FY_MONTHS:
            g, rt, cn, ib = base["GTS"][s][m], base["Returns"][s][m], base["Credit Notes"][s][m], base["Inter Sales"][s][m]
            rr = ret[s][m] or z()
            ns = {k: g[k] - rt[k] - cn[k] for k in keys}
            df = {k: rr[k] - ns[k] for k in keys}
            row = {"GTS": g, "Returns": rt, "Credit Notes": cn, "Net Sales": ns, "Return": rr, "Difference": df,
                   "Inter Sales": ib, "Net Difference": {k: df[k] - ib[k] for k in keys}}
            for ln in lines:
                vals[ln][s][m] = {k: round(row[ln][k], 2) for k in keys}
    for ln in lines:                                                      # FY and all-registration totals
        for s in regs:
            vals[ln][s]["FY"] = {k: round(sum(vals[ln][s][m][k] for m in _FY_MONTHS), 2) for k in keys}
        vals[ln]["Total"] = {m: {k: round(sum(vals[ln][s][m][k] for s in regs), 2) for k in keys}
                             for m in _FY_MONTHS + ["FY"]}
    sources = {s: {m: (res["comp"][s][m].get("doc") or "") for m in _FY_MONTHS} for s in regs}
    return {
        "lines": lines,
        "registrations": [{"key": s, "label": labels.get(s, s), "gstin": res["gstin_of"].get(s, "")} for s in regs],
        "months": list(_FY_MONTHS),
        "values": vals,
        "return_source": sources,
        "fy": res["fy"],
    }
