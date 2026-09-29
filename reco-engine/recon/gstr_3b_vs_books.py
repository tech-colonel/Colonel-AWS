"""GSTR-3B vs Books — GST Summary of the Sales Register + GSTR-3B, every state, every month.

Books follow the accountant's GST Summary working (see gst_books_summary):
state from the tax columns, Interbranch shown separately, Returns deducted, the
register tied back to its own Grand Total. GSTR-3B table 3.1 — the GST-portal PDF
(read with the GSTR-3B Tally Entry agent's parser) or an OCTA GSTR3B sheet — is
compared per registration x month against Books "Total incl. Interbranch" (Sales −
Returns + Interbranch), which is what 3.1(a)+(b) declares.
"""
from __future__ import annotations

import logging
from datetime import date, datetime

from .core import round_money
from .gst_books_summary import AMT, read_books, summarise, tie_out, _canon_state
from .gstr_1_multistate import _amt, read_return_files, state_of_gstin
from .gstr_1_vs_books import _FY_MONTHS, _f

logger = logging.getLogger(__name__)


def _month_end(fy_start: int, month: str) -> date:
    import calendar
    mo = datetime.strptime(month, "%B").month
    yr = fy_start if mo >= 4 else fy_start + 1
    return date(yr, mo, calendar.monthrange(yr, mo)[1])


def _tax(d: dict) -> float:
    return round_money(_f(d.get("cgst", 0)) + _f(d.get("sgst", 0)) + _f(d.get("igst", 0)))


def run_gstr3b_vs_books(tally_items: list[dict], cn_items: list[dict] | None,
                        return_items: list[dict], tolerance: float) -> dict:
    books = read_books(tally_items, cn_items)
    summ = summarise(books)
    tie = tie_out(books, summ)

    dates = [r["date"] for r in books["rows"] if r["date"]]
    first = min(dates) if dates else datetime.today()
    fy_start = first.year if first.month >= 4 else first.year - 1
    fy = f"{fy_start}-{str(fy_start + 1)[2:]}"

    # ---- GSTR-3B, per canonical state x month
    ret = read_return_files(return_items) if return_items else None
    g3b: dict[str, dict] = {}
    gstin_of: dict[str, str] = {}
    detail = []
    if ret:
        for gstin, per in ret["g3b_pdf"].items():
            st = _canon_state(state_of_gstin(gstin)) or state_of_gstin(gstin)
            gstin_of[st] = gstin
            for p, v in per.items():
                rows = v["rows"]
                g3b.setdefault(st, {})[v["month"]] = {
                    "amt": _amt(v["outward"]), "source": "Portal PDF", "file": v.get("_file"),
                    "c": round_money(_f((rows.get("c") or {}).get("taxable", 0))),
                    "e": round_money(_f((rows.get("e") or {}).get("taxable", 0))),
                    "arn_date": v.get("filing_date", ""), "ok": v["parsed_ok"], "period": p, "rows": rows,
                    "doc": "GSTR-3B", "nil": False,
                }
        for gstin, per in ret["g3b_octa"].items():
            st = _canon_state(state_of_gstin(gstin)) or state_of_gstin(gstin)
            gstin_of.setdefault(st, gstin)
            for month, amt in per.items():
                g3b.setdefault(st, {}).setdefault(month, {
                    "amt": _amt(amt), "source": "OCTA Excel", "file": "", "c": None, "e": None,
                    "arn_date": "", "ok": True, "period": "", "rows": {}, "doc": "GSTR-3B", "nil": False})
        # GSTR-1 and GSTR-3B declare the same outward supplies in a different layout, so
        # a month with no GSTR-3B is compared with that month's GSTR-1 instead (portal
        # PDF "Total Liability", or the OCTA invoice rows added up) — labelled as such.
        # A GSTR-3B for the same registration x month always takes precedence.
        from .gstr_1_vs_books import aggregate_gstr1_monthly
        for gstin, per in ret["g1_pdf"].items():
            st = _canon_state(state_of_gstin(gstin)) or state_of_gstin(gstin)
            gstin_of.setdefault(st, gstin)
            for p, v in per.items():
                if v["month"] in g3b.get(st, {}):
                    continue
                g3b.setdefault(st, {})[v["month"]] = {
                    "amt": _amt(v["total"]), "source": "Portal PDF", "file": v.get("_file"), "c": None, "e": None,
                    "arn_date": "", "ok": v["parsed_ok"], "period": p, "rows": {}, "doc": "GSTR-1",
                    "nil": v["nil_filed"]}
        for gstin, df in ret["octa"].items():
            st = _canon_state(state_of_gstin(gstin)) or state_of_gstin(gstin)
            gstin_of.setdefault(st, gstin)
            for month, amt in aggregate_gstr1_monthly(df).items():
                if month in _FY_MONTHS and month not in g3b.get(st, {}):
                    g3b.setdefault(st, {})[month] = {
                        "amt": _amt(amt), "source": "OCTA Excel", "file": "", "c": None, "e": None,
                        "arn_date": "", "ok": True, "period": "", "rows": {}, "doc": "GSTR-1", "nil": False}
        for st, months in g3b.items():
            for m, v in months.items():
                a = v["rows"].get("a") or {}
                detail.append({
                    "Registration": summ["labels"].get(st, st), "GSTIN": gstin_of.get(st, ""), "Month": m,
                    "Compared with": v["doc"], "Format": v["source"],
                    "Taxable Value": a.get("taxable", v["amt"]["taxable"]) if v["doc"] == "GSTR-3B" else v["amt"]["taxable"],
                    "IGST": a.get("igst", v["amt"]["igst"]) if v["doc"] == "GSTR-3B" else v["amt"]["igst"],
                    "CGST": a.get("cgst", v["amt"]["cgst"]) if v["doc"] == "GSTR-3B" else v["amt"]["cgst"],
                    "SGST": a.get("sgst", v["amt"]["sgst"]) if v["doc"] == "GSTR-3B" else v["amt"]["sgst"],
                    "3B 3.1(b) Zero-rated": (v["rows"].get("b") or {}).get("taxable"),
                    "3B 3.1(c) Nil / Exempt": v["c"], "3B 3.1(e) Non-GST": v["e"],
                    "Nil return": "Yes" if v.get("nil") else "",
                    "ARN date": v["arn_date"], "File": v["file"], "Read OK": "Yes" if v["ok"] else "Check",
                })
        detail.sort(key=lambda r: (r["Registration"], _FY_MONTHS.index(r["Month"])))

    # ---- Compare: Books total incl. Interbranch vs 3B, per state x month
    states = list(summ["states"])
    for st in g3b:
        if st not in states:
            states.append(st)
            summ["labels"].setdefault(st, st)
    comp = {}
    for st in states:
        comp[st] = {}
        for m in _FY_MONTHS:
            b = (summ["monthly"]["total"].get(st) or {}).get(m) or {k: 0.0 for k in AMT}
            g = (g3b.get(st) or {}).get(m)
            has = g is not None
            ga = ({"net": g["amt"]["taxable"], "cgst": g["amt"]["cgst"], "sgst": g["amt"]["sgst"],
                   "igst": g["amt"]["igst"]} if has else None)
            d = {k: round_money(b[k] - ga[k]) for k in AMT} if has else None
            doc = g["doc"] if has else ""
            if not has:
                status = "Return not uploaded" if any(abs(b[k]) > 0 for k in AMT) else ""
            elif all(abs(d[k]) <= tolerance for k in AMT):
                status = "Matched" + ("" if doc == "GSTR-3B" else " (vs GSTR-1)")
            elif d["net"] > tolerance or _tax(d) > tolerance:
                status = f"Short in {doc}"
            else:
                status = f"Excess in {doc}"
            if has and g.get("nil") and status.startswith("Short"):
                status = f"{doc} filed Nil"
            comp[st][m] = {"books": b, "g3b": ga, "diff": d, "status": status,
                           "source": g["source"] if has else "", "doc": doc}
    # Net Position (Sales excl. Interbranch − Returns) vs 3B — the table right after
    # Section 4. Interbranch supplies sit inside 3B too, so a registration that made
    # them shows that amount as a difference here (Section 8 compares incl. them).
    comp_net = {}
    for st in states:
        comp_net[st] = {}
        for m in _FY_MONTHS:
            b = (summ["monthly"]["net"].get(st) or {}).get(m) or {k: 0.0 for k in AMT}
            ib = (summ["monthly"]["interbranch"].get(st) or {}).get(m) or {k: 0.0 for k in AMT}
            g = comp[st][m]["g3b"]
            doc = comp[st][m]["doc"]
            vs = "" if doc == "GSTR-3B" else " (vs GSTR-1)"
            d = {k: round_money(b[k] - g[k]) for k in AMT} if g is not None else None
            if g is None:
                status = "Return not uploaded" if any(abs(b[k]) > 0 for k in AMT) else ""
            elif all(abs(d[k]) <= tolerance for k in AMT):
                status = "Matched" + vs
            elif d["net"] > tolerance or _tax(d) > tolerance:
                status = f"Short in {doc}"
            else:
                status = f"Excess in {doc}"
            if g is not None and not status.startswith("Matched") and any(abs(ib[k]) > 0 for k in AMT) and \
                    all(abs(d[k] + ib[k]) <= tolerance for k in AMT):
                status = "Matched incl. Interbranch" + vs
            if g is not None and (g3b.get(st) or {}).get(m, {}).get("nil") and status.startswith("Short"):
                status = f"{doc} filed Nil"
            comp_net[st][m] = {"books": b, "g3b": g, "diff": d, "status": status, "interbranch": ib, "doc": doc}
    comp_net_annual = {}
    for st in states:
        months = [m for m in _FY_MONTHS if comp_net[st][m]["g3b"] is not None]
        b = {k: round_money(sum(comp_net[st][m]["books"][k] for m in months)) for k in AMT}
        g = {k: round_money(sum(comp_net[st][m]["g3b"][k] for m in months)) for k in AMT}
        comp_net_annual[st] = {"months": len(months), "books": b, "g3b": g,
                               "diff": {k: round_money(b[k] - g[k]) for k in AMT},
                               "matched": sum(1 for m in months if comp_net[st][m]["status"].startswith("Matched")),
                               "vs_gstr1": sum(1 for m in months if comp_net[st][m]["doc"] == "GSTR-1")}

    comp_annual = {}
    for st in states:
        months = [m for m in _FY_MONTHS if comp[st][m]["g3b"] is not None]
        b = {k: round_money(sum(comp[st][m]["books"][k] for m in months)) for k in AMT}
        g = {k: round_money(sum(comp[st][m]["g3b"][k] for m in months)) for k in AMT}
        comp_annual[st] = {"months": len(months), "books": b, "g3b": g,
                           "diff": {k: round_money(b[k] - g[k]) for k in AMT},
                           "matched": sum(1 for m in months if comp[st][m]["status"].startswith("Matched")),
                           "vs_gstr1": sum(1 for m in months if comp[st][m]["doc"] == "GSTR-1")}

    # ---- Checks
    checks = [{"check": "Every register row is Sales or Returns and sits in a state",
               "expected": tie["n_rows"], "actual": tie["n_sales"] + tie["n_returns"],
               "ok": tie["n_rows"] == tie["n_sales"] + tie["n_returns"]}]
    for b in tie["blocks"]:
        if b["variance"] is not None:
            checks.append({"check": f"Ties to its own total row — {b['label']}",
                           "expected": b["grand_total"]["net"], "actual": b["compared"]["net"],
                           "ok": all(abs(v) < 0.01 for v in b["variance"].values()),
                           "note": b["basis"] + (f"; {len(b['bridges'])} bridging item(s)" if b["bridges"] else "")})
    if ret:
        p3 = [p for g in ret["g3b_pdf"].values() for p in g.values()]
        if p3:
            bad = [p for p in p3 if not p["parsed_ok"]]
            checks.append({"check": "Every GSTR-3B PDF: table 3.1(a) found", "expected": len(p3),
                           "actual": len(p3) - len(bad), "ok": not bad})
        p1 = [p for g in ret["g1_pdf"].values() for p in g.values()]
        if p1:
            bad1 = [p for p in p1 if not p["parsed_ok"]]
            checks.append({"check": "Every GSTR-1 PDF: table totals = the form's own Total Liability line",
                           "expected": len(p1), "actual": len(p1) - len(bad1), "ok": not bad1})

    mismatch_rows = [r for r in books["rows"] if r["state_mismatch"]]
    interbranch_rows = [r for r in books["rows"] if r["interbranch"]]
    notes = [
        "Sales vs Returns follows the \"Nature\" column (else a Credit Note voucher is a return). Returns are shown "
        "as positive figures and deducted in Section 4.",
        f"Interbranch Services ({len(interbranch_rows)} row(s)) are excluded from the state-wise Sales in Section 1 "
        "and shown on their own in Section 2. Section 5 adds them back so the total reconciles to the register.",
        "State: each row is assigned to the state of its tax columns (state-tagged ledgers; untagged ledgers = home "
        "registration). Rows with no tax, and blocks with plain CGST/SGST/IGST columns, use the State field. "
        + (f"{len(mismatch_rows)} row(s) where the State field differs — register rows "
           + ", ".join(str(r['row']) for r in mismatch_rows[:40]) + (" …" if len(mismatch_rows) > 40 else "")
           + ". Please confirm." if mismatch_rows else "No row's State field disagrees with its tax columns."),
        "Each registration-month is compared with its GSTR-3B (table 3.1(a)+(b)); where no GSTR-3B was uploaded, "
        "with that month's GSTR-1 (portal PDF Total Liability, or the OCTA rows added up) — GSTR-1 and GSTR-3B "
        "declare the same outward supplies. Section 4A compares Net Position, Section 8 the total incl. "
        "Interbranch. Months with neither return are listed, not compared.",
    ]
    if len(tie["blocks"]) > 1:
        notes.append(f"The upload held {len(tie['blocks'])} blocks, read separately (each with its own columns) "
                     "and added together: " + "; ".join(f"{b['label']} — {b['rows']} rows" for b in tie["blocks"]) + ".")

    return {
        "fy": fy, "fy_start": fy_start, "company": _company_title(tally_items), "books": books, "summ": summ,
        "tie": tie, "states": states, "g3b": g3b, "gstin_of": gstin_of, "comp": comp, "comp_annual": comp_annual,
        "comp_net": comp_net, "comp_net_annual": comp_net_annual,
        "detail": detail, "checks": checks, "notes": notes, "mismatch_rows": mismatch_rows,
        "files": ret["files"] if ret else [], "warnings": (ret["warnings"] if ret else []),
        "month_ends": {m: _month_end(fy_start, m) for m in _FY_MONTHS},
    }


def _company_title(items) -> str:
    """First non-empty text above the first header — Tally prints the company name there."""
    from io import BytesIO
    from openpyxl import load_workbook
    for it in items:
        try:
            ws = load_workbook(BytesIO(it["content"]), read_only=True, data_only=True).worksheets[0]
            for r in ws.iter_rows(min_row=1, max_row=6, values_only=True):
                v = next((c for c in r if c not in (None, "")), None)
                if isinstance(v, str) and v.strip():
                    return v.strip()
        except Exception:
            continue
    return ""
