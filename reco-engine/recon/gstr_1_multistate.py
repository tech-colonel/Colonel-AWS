"""GSTR-1 vs Books — combined Sales Register, every state, every month.

A firm registered in several states files one GSTR-1 and one GSTR-3B per GSTIN per
month, but keeps ONE Sales Register in Tally covering every registration. The
single-state engine (`gstr_1_vs_books`) takes one register and one GSTR-1 file,
so this module sits in front of it:

  1. read every return file — GSTR-1 as OCTA Excel (invoice level) or as the GST
     portal PDF (table totals), GSTR-3B as the OCTA sheet or the portal PDF — work
     out the GSTIN and month of each, and drop a registration x month that a later
     file repeats (a yearly file plus monthly files must not double count);
  2. split the combined Sales Register by registration (`assign_books_states`);
  3. run the EXISTING single-state steps, unchanged, once per registration;
  4. add them up into an "All States" view;
  5. flag an invoice booked under one state but filed in another state's GSTR-1
     (Remark 3). Remark 1 is never changed.

No reconciliation rule lives here — rows are only routed to the existing
functions. Every Books row lands in exactly one registration or in "Unassigned",
and `_integrity_checks` proves nothing was dropped.
"""
from __future__ import annotations

import hashlib
import logging
import re
from collections import Counter, defaultdict
from datetime import datetime
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from .core import round_money
from .gstr1_portal_pdf import parse_gstr1_portal_pdf, parse_gstr3b_outward_pdf
from .gstr_1_vs_books import (
    _FY_MONTHS,
    _add_amounts,
    _classify_tally_ledger,
    _col_val,
    _f,
    _find_col,
    _gstr1_b2b_invoice_set,
    _norm_inv,
    _norm_month,
    _sorted_months,
    _zero_amounts,
    aggregate_books_monthly,
    aggregate_gstr1_monthly,
    build_monthly_comparison,
    build_summary,
    df_to_records,
    extract_gstr3b_monthly,
    gstr1_month_basis,
    merge_credit_notes,
    read_octa_excel,
    reconcile_b2b_new,
    reconcile_b2c_new,
)
from .gstr_2b_books import GST_STATE_CODES, _state_from_voucher_type
from .parsers import normalize_header

logger = logging.getLogger(__name__)

UNASSIGNED = "Unassigned"
# Books rows that clearly belong to a registration of the firm whose return was not
# uploaded in this run ("NOREG:Maharashtra"). Kept apart — never lumped into another
# state — and shown, but not compared with any return.
NOREG = "NOREG:"
_AMOUNT_KEYS = ("taxable", "igst", "cgst", "sgst")
_GSTIN_RE = re.compile(r"[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][0-9A-Z]Z[0-9A-Z]")

# Short state tags Tally users put in per-registration ledger / column values
# ("Output CGST 9% DL", "Sales KAR", "Output IGST BLR 5%"). Trusted only when the
# tag resolves to a registration actually uploaded.
_STATE_TAGS = {
    "DL": "07", "DEL": "07", "HR": "06", "HAR": "06", "KA": "29", "KAR": "29",
    "BLR": "29", "BNG": "29", "MH": "27", "MAH": "27", "TN": "33", "UP": "09",
    "GJ": "24", "GUJ": "24", "WB": "19", "RJ": "08", "RAJ": "08", "TS": "36",
    "TG": "36", "AP": "37", "KL": "32", "PB": "03", "MP": "23", "BR": "10",
    "OD": "21", "OR": "21", "AS": "18", "JH": "20", "CG": "22", "UK": "05",
    "HP": "02", "JK": "01", "GA": "30", "CH": "04",
}
# Warehouse / branch cities people write instead of the state name.
_CITY_STATE = {
    "bangalore": "Karnataka", "bengaluru": "Karnataka", "mysore": "Karnataka",
    "mumbai": "Maharashtra", "pune": "Maharashtra", "vasai": "Maharashtra",
    "bhiwandi": "Maharashtra", "thane": "Maharashtra", "nagpur": "Maharashtra",
    "gurgaon": "Haryana", "gurugram": "Haryana", "faridabad": "Haryana", "manesar": "Haryana",
    "noida": "Uttar Pradesh", "lucknow": "Uttar Pradesh", "new delhi": "Delhi",
    "chennai": "Tamil Nadu", "hyderabad": "Telangana", "kolkata": "West Bengal",
    "ahmedabad": "Gujarat", "surat": "Gujarat", "jaipur": "Rajasthan",
    "kochi": "Kerala", "indore": "Madhya Pradesh", "bhopal": "Madhya Pradesh",
    "guwahati": "Assam", "patna": "Bihar", "ludhiana": "Punjab",
}
# Register columns that name the seller's own registration outright.
_REG_COL_EXPLICIT = ["Company GSTIN", "Registration GSTIN", "Our GSTIN", "Seller GSTIN",
                     "Registration", "Registration State", "Company State"]
# Columns that MAY hold the registration — or may be the customer's place of
# supply. Used only when nearly every value resolves to an uploaded registration.
_REG_COL_MAYBE = ["State", "States", "Branch", "Location", "Godown", "Warehouse", "Unit"]
# A "Nature"-style column that marks sales returns inside the Sales Register.
_NATURE_COLS = ["Nature", "Nature of Transaction", "Transaction Type", "Txn Type", "Entry Type"]
_RETURN_RE = re.compile(r"return|credit\s*note|\brto\b", re.I)


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

def _clean_gstin(v) -> str:
    s = re.sub(r"[^A-Z0-9]", "", str(v or "").upper())
    return s if len(s) == 15 and _GSTIN_RE.fullmatch(s) else ""


def _base_state(name: str) -> str:
    """'Andhra Pradesh (New)' and 'Andhra Pradesh' are one place for matching."""
    return re.sub(r"\s*\(new\)\s*$", "", str(name or "").strip(), flags=re.I).lower()


def state_of_gstin(gstin: str) -> str:
    return GST_STATE_CODES.get((gstin or "")[:2], "") if gstin else ""


def is_real_reg(reg: str) -> bool:
    return bool(reg) and reg != UNASSIGNED and not reg.startswith(NOREG)


def reg_label(reg: str) -> str:
    if reg == UNASSIGNED:
        return UNASSIGNED
    if reg.startswith(NOREG):
        return f"{reg[len(NOREG):]} (no return uploaded)"
    return state_of_gstin(reg) or reg


def reg_display(reg: str) -> str:
    return f"{reg_label(reg)} ({reg})" if is_real_reg(reg) else reg_label(reg)


def _period_key(v) -> str:
    """Tax Period → 'YYYY-MM' (the year kept, unlike the month-name buckets)."""
    if v is None:
        return ""
    try:
        if pd.isnull(v):
            return ""
    except Exception:
        pass
    if isinstance(v, (datetime, pd.Timestamp)):
        return v.strftime("%Y-%m")
    text = str(v).strip()
    try:
        return datetime.fromisoformat(text[:10]).strftime("%Y-%m")
    except Exception:
        pass
    m = re.match(r"^(\d{2})(\d{4})$", text)              # portal style 042025
    if m:
        return f"{m.group(2)}-{m.group(1)}"
    for fmt in ("%b %Y", "%B %Y", "%b-%Y", "%B-%Y", "%b-%y", "%m-%Y", "%m/%Y"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m")
        except Exception:
            continue
    return text.lower()


def _fy_of(period_key: str) -> str:
    m = re.match(r"^(\d{4})-(\d{2})$", period_key or "")
    if not m:
        return ""
    y, mo = int(m.group(1)), int(m.group(2))
    start = y if mo >= 4 else y - 1
    return f"{start}-{str(start + 1)[2:]}"


def _amt(d: dict | None) -> dict:
    d = d or {}
    return {k: round_money(_f(d.get(k, 0))) for k in _AMOUNT_KEYS}


def _sum_months(monthly: dict) -> dict:
    tot = _zero_amounts()
    for m, amt in (monthly or {}).items():
        if m in _FY_MONTHS:
            for k in _AMOUNT_KEYS:
                tot[k] += _f(amt.get(k, 0))
    return {k: round_money(v) for k, v in tot.items()}


# ---------------------------------------------------------------------------
# Books — register-level preparation
# ---------------------------------------------------------------------------

def normalise_returns(df: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    """
    Some Sales Registers carry sales returns as ordinary rows with POSITIVE amounts
    and a 'Nature' column saying "Returns". Left as they are, every return adds to
    turnover instead of reducing it. Rows marked as a return whose amount is still
    positive are sign-flipped (ledger columns, the derived totals, Gross Total);
    returns already entered as negatives are left alone.
    """
    info = {"column": None, "flipped": 0, "already_negative": 0}
    col = _find_col(df, _NATURE_COLS)
    if col is None or df.empty:
        return df, info
    info["column"] = col
    ts_col = _find_col(df, ["Total Sales", "Taxable Value", "Taxable Amount"])
    amount_cols = [c for c in df.columns if _classify_tally_ledger(c) and c != col]
    for extra in ("Gross Total", "Value", "Total Sales", "Total IGST", "Total CGST", "Total SGST", "Total Cess"):
        c = _find_col(df, [extra])
        if c and c not in amount_cols:
            amount_cols.append(c)
    df = df.copy()
    flip_idx = []
    for idx, row in df.iterrows():
        if not _RETURN_RE.search(str(_col_val(row, col, "") or "")):
            continue
        total = _f(_col_val(row, ts_col, 0)) if ts_col else 0.0
        if total > 0:
            flip_idx.append(idx)
        elif total < 0:
            info["already_negative"] += 1
    for c in amount_cols:
        df.loc[flip_idx, c] = [(-_f(v) if _f(v) else v) for v in df.loc[flip_idx, c]]
    info["flipped"] = len(flip_idx)
    if flip_idx:
        logger.info("Sales Register: %d return row(s) in '%s' carried positive amounts — "
                    "sign-flipped so they reduce turnover", len(flip_idx), col)
    return df, info


def register_grand_total(file_info: dict) -> dict | None:
    """Tally's own Grand Total row of the register (dropped by the reader), summed by
    kind — the figure the Books side must tie to. None when the export has none."""
    try:
        wb = load_workbook(BytesIO(file_info["content"]), data_only=True, read_only=True)
    except Exception:
        return None
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        header = None
        for r in rows[:20]:
            nh = [normalize_header(v) for v in r if v]
            if "particulars" in nh or "date" in nh:
                header = [str(v or "").strip() for v in r]
                break
        if not header:
            continue
        for r in rows:
            if any(str(v or "").strip().lower() == "grand total" for v in r[:6]):
                tot = _zero_amounts()
                for h, v in zip(header, r):
                    kind = _classify_tally_ledger(h)
                    if kind in tot and normalize_header(h) not in {normalize_header(n) for n in _NATURE_COLS}:
                        tot[kind] += _f(v)
                return {k: round_money(v) for k, v in tot.items()}
    return None


_ALL_STATE_NAMES = {_base_state(n): n for n in GST_STATE_CODES.values()}


def _state_name_of_text(text: str) -> str:
    """'Karnataka' / 'Bangalore' / '29-Karnataka' / 'KAR' / a GSTIN → canonical state name, or ''."""
    raw = str(text or "").strip()
    if not raw:
        return ""
    g = _clean_gstin(raw)
    if g:
        return state_of_gstin(g)
    t = re.sub(r"^\d{1,2}\s*[-–]\s*", "", raw)
    if _base_state(t) in _ALL_STATE_NAMES:
        return _ALL_STATE_NAMES[_base_state(t)]
    low = t.lower()
    if low in _CITY_STATE:
        return _CITY_STATE[low]
    if t.upper() in _STATE_TAGS:
        return GST_STATE_CODES.get(_STATE_TAGS[t.upper()], "")
    named = _state_from_voucher_type(t)
    if named:
        return named
    for city, st in _CITY_STATE.items():
        if city in low:
            return st
    return ""


def _resolve_state_text(text: str, regs: list[str]) -> str:
    """The uploaded registration the text names, else NOREG:<State> when it names a
    state with no return in this run, else ''."""
    got = _state_text_to_reg(text, regs)
    if got:
        return got
    name = _state_name_of_text(text)
    return f"{NOREG}{name}" if name else ""


def _state_text_to_reg(text: str, regs: list[str]) -> str:
    """'Karnataka' / 'Bangalore' / '29-Karnataka' / 'KAR' / a GSTIN → the uploaded registration."""
    raw = str(text or "").strip()
    if not raw:
        return ""
    g = _clean_gstin(raw)
    if g:
        return g if g in regs else ""
    by_state = defaultdict(list)
    for r in regs:
        by_state[_base_state(state_of_gstin(r))].append(r)

    def pick(name):
        hits = by_state.get(_base_state(name), [])
        return hits[0] if len(hits) == 1 else ""

    t = re.sub(r"^\d{1,2}\s*[-–]\s*", "", raw)
    got = pick(t)
    if got:
        return got
    low = t.lower()
    if low in _CITY_STATE:
        return pick(_CITY_STATE[low])
    if t.upper() in _STATE_TAGS:
        code = _STATE_TAGS[t.upper()]
        hits = [r for r in regs if r[:2] == code or (code == "37" and r[:2] == "28")]
        return hits[0] if len(hits) == 1 else ""
    named = _state_from_voucher_type(t)
    if named:
        return pick(named)
    for city, st in _CITY_STATE.items():
        if city in low:
            return pick(st)
    return ""


def register_row_checks(df: pd.DataFrame) -> list[dict]:
    """Rows whose ledger columns (+ Round Off) do not add up to their Gross Total.

    Tally exports Dr/Cr as bare numbers, so a ledger that was debited on a sales
    voucher arrives with the wrong sign; the row then no longer foots. Reported only
    — the amounts are used exactly as exported."""
    gross_col = _find_col(df, ["Gross Total"])
    if gross_col is None or df.empty:
        return []
    ledger_cols = [c for c in df.columns if _classify_tally_ledger(c)
                   and normalize_header(c) not in {normalize_header(n) for n in _NATURE_COLS}]
    ro = _find_col(df, ["Round Off", "Rounding Off", "Round Off (Sales)"])
    if ro:
        ledger_cols.append(ro)
    date_col = _find_col(df, ["Date", "Voucher Date", "Invoice Date"])
    part_col = _find_col(df, ["Particulars", "Party Name", "Buyer"])
    narr_col = _find_col(df, ["Narration"])
    out = []
    for _, row in df.iterrows():
        gross = _f(_col_val(row, gross_col, 0))
        total = sum(_f(_col_val(row, c, 0)) for c in ledger_cols)
        if abs(abs(gross) - abs(total)) > 1.0:
            out.append({
                "Date": _col_val(row, date_col, ""), "Particulars": _col_val(row, part_col, ""),
                "Narration": _col_val(row, narr_col, ""), "Gross Total": round_money(abs(gross)),
                "Sum of ledger columns": round_money(abs(total)),
                "Difference": round_money(abs(total) - abs(gross)),
            })
    return out


def _pick_registration_column(df: pd.DataFrame, regs: list[str]) -> str | None:
    col = _find_col(df, _REG_COL_EXPLICIT)
    if col:
        return col
    for cand in _REG_COL_MAYBE:
        col = _find_col(df, [cand])
        if not col:
            continue
        vals = [str(v).strip() for v in df[col] if v is not None and str(v).strip() and str(v).lower() != "nan"]
        if not vals:
            continue
        distinct = set(vals)
        states_named = {v for v in distinct if _state_name_of_text(v)}
        share = sum(1 for v in vals if v in states_named) / len(vals)
        hits_upload = any(_state_text_to_reg(v, regs) for v in distinct)
        # A registration column names only the firm's own few registrations (even
        # those with no return uploaded in this run); a place-of-supply column names
        # customers' states — many of them.
        if share >= 0.95 and hits_upload and len(distinct) <= max(len(regs) * 3, 6):
            return col
    return None


def assign_books_states(reco_df: pd.DataFrame, gstr1_by_reg: dict[str, pd.DataFrame],
                        regs: list[str]) -> tuple[pd.Series, pd.Series]:
    """
    Registration GSTIN + reason for every Books row. Evidence, in order:

      1. a registration column (Company GSTIN / Registration / or a State/Branch
         column whose values are only the firm's own registrations);
      2. the Tally Voucher Type ("Sales Karnataka", "Credit Note Delhi");
      3. the invoice number is in exactly ONE registration's GSTR-1 (OCTA);
      4. every non-zero state-tagged ledger ("Output CGST BLR 9%") points at one
         registration;
      5. intra-state supply (CGST/SGST, no IGST) to a place of supply that is one
         of the registrations — by law only that registration can make it;
      6. only one registration uploaded;
      otherwise "Unassigned" (kept, listed, still counted in All States).
    """
    reg_col = _pick_registration_column(reco_df, regs)
    vt_col = _find_col(reco_df, ["Voucher Type", "Vch Type", "Voucher Type Name"])
    inv_col = _find_col(reco_df, ["Voucher No.", "Voucher No", "Invoice No", "Doc No", "Bill No"])
    ref_col = _find_col(reco_df, ["Voucher Ref. No.", "Voucher Ref No", "Ref No"])
    igst_col = _find_col(reco_df, ["Total IGST", "IGST"])
    cgst_col = _find_col(reco_df, ["Total CGST", "CGST"])
    sgst_col = _find_col(reco_df, ["Total SGST", "SGST"])
    pos_col = _find_col(reco_df, ["States", "Place of Supply", "State"])
    gstin_col = _find_col(reco_df, ["GSTIN", "GSTIN/UIN", "Buyer GSTIN"])

    inv_regs: dict[str, set] = defaultdict(set)
    for g, gdf in gstr1_by_reg.items():
        inv_col_g = _find_col(gdf, ["Doc No", "Invoice Number", "Invoice No"])
        if inv_col_g:
            for v in gdf[inv_col_g]:
                n = _norm_inv(v)
                if n:
                    inv_regs[n].add(g)

    ledger_reg: dict[str, str] = {}
    for col in reco_df.columns:
        if not _classify_tally_ledger(col) or col == reg_col:
            continue
        text = str(col)
        reg = ""
        for tok in [t for t in re.split(r"[\s\-_/()%.\d]+", text) if t]:
            if tok.upper() in _STATE_TAGS or tok.lower() in _CITY_STATE:
                reg = _state_text_to_reg(tok, regs)
                if reg:
                    break
        if not reg:
            named = _state_from_voucher_type(text)
            reg = _state_text_to_reg(named, regs) if named else ""
        if reg:
            ledger_reg[col] = reg

    out_reg, out_why = [], []
    for _, row in reco_df.iterrows():
        reg, why = "", ""
        if reg_col:
            raw = str(_col_val(row, reg_col, "") or "").strip()
            reg = _resolve_state_text(raw, regs) if raw else ""
            if reg:
                why = f"'{reg_col}' column = {raw}" + ("" if is_real_reg(reg) else " — no return uploaded for it")
            elif raw:
                why = f"'{reg_col}' = {raw}, which is not a state"
        if not reg and vt_col:
            vt = str(_col_val(row, vt_col, "") or "").strip()
            st = _state_from_voucher_type(vt)
            if st:
                reg = _resolve_state_text(st, regs)
                why = f"Voucher Type '{vt}'" + ("" if is_real_reg(reg) else " — no return uploaded for it")
        if not reg and inv_regs:
            hits = set()
            for c in (inv_col, ref_col):
                if c:
                    n = _norm_inv(_col_val(row, c, ""))
                    if n and n in inv_regs:
                        hits |= inv_regs[n]
            if len(hits) == 1:
                reg, why = next(iter(hits)), "Invoice number found in this registration's GSTR-1"
        if not reg and ledger_reg:
            tagged = {ledger_reg[c] for c in ledger_reg if abs(_f(_col_val(row, c, 0))) > 0}
            if len(tagged) == 1:
                reg, why = next(iter(tagged)), "State-tagged ledger column"
        if not reg:
            igst = abs(_f(_col_val(row, igst_col, 0)))
            intra = abs(_f(_col_val(row, cgst_col, 0))) + abs(_f(_col_val(row, sgst_col, 0)))
            if igst == 0 and intra > 0:
                pos = str(_col_val(row, pos_col, "") or "").strip() if pos_col else ""
                if not pos and gstin_col:
                    pos = state_of_gstin(_clean_gstin(_col_val(row, gstin_col, "")))
                got = _state_text_to_reg(pos, regs) if pos else ""
                if got:
                    reg, why = got, "Intra-state supply (CGST/SGST) in this state"
        if not reg and len(regs) == 1:
            reg, why = regs[0], "Only one registration uploaded"
        if not reg:
            reg = UNASSIGNED
            why = why or "No registration column, voucher-type state, invoice match, state ledger or intra-state signal"
        out_reg.append(reg)
        out_why.append(why)
    return pd.Series(out_reg, index=reco_df.index), pd.Series(out_why, index=reco_df.index)


def books_split_available(reco_df: pd.DataFrame) -> bool:
    """B2B vs B2C in Books needs a Category column or buyer GSTINs; a consolidated
    register (one entry per channel per month) has neither."""
    if _find_col(reco_df, ["Category", "Catogary", "Cat"]):
        return True
    gc = _find_col(reco_df, ["GSTIN", "GSTIN/UIN", "Buyer GSTIN", "GST No"])
    return bool(gc) and any(_clean_gstin(v) for v in reco_df[gc])


# ---------------------------------------------------------------------------
# Returns — read every file, bucket by registration and period
# ---------------------------------------------------------------------------

def read_return_files(items: list[dict]) -> dict:
    files: list[dict] = []
    warnings: list[str] = []
    seen_hash: dict[str, str] = {}
    octa_seen: dict[tuple, str] = {}        # (gstin, period) → file
    octa_frames: dict[str, list] = defaultdict(list)
    raw_g1: list[pd.DataFrame] = []
    raw_3b: list[pd.DataFrame] = []
    raw_2b: list[pd.DataFrame] = []
    g1_pdf: dict[str, dict] = defaultdict(dict)        # gstin → period → parsed
    g3b_pdf: dict[str, dict] = defaultdict(dict)
    g3b_octa: dict[str, dict] = defaultdict(dict)       # gstin → month → amounts
    pending_3b_sheets: list[tuple] = []
    undetected: list[tuple] = []
    octa_pending: list[tuple] = []

    for i, item in enumerate(items):
        fname = item.get("filename") or f"file {i + 1}"
        content = item.get("content") or b""
        info = {"file": fname, "kind": "", "gstin": "", "state": "", "periods": [],
                "rows": 0, "status": "", "note": ""}
        files.append(info)
        if not content:
            info.update(status="Skipped", note="Empty file")
            continue
        digest = hashlib.sha1(content).hexdigest()
        if digest in seen_hash:
            info.update(status="Skipped", note=f"Same file as {seen_hash[digest]}")
            continue
        seen_hash[digest] = fname

        if content[:4] == b"%PDF" or fname.lower().endswith(".pdf"):
            p1 = parse_gstr1_portal_pdf(content)
            if p1:
                g, per = p1["gstin"], p1["period"]
                info.update(kind="GSTR-1 (portal PDF)", gstin=g, state=state_of_gstin(g), periods=[per])
                if per in g1_pdf[g]:
                    info.update(status="Skipped", note=f"{per} already read from {g1_pdf[g][per]['_file']}")
                    continue
                p1["_file"] = fname
                g1_pdf[g][per] = p1
                info["status"] = "Used"
                notes = []
                if p1["nil_filed"]:
                    notes.append("Nil return filed")
                if not p1["parsed_ok"]:
                    notes.append("Table totals do not add up to the form's Total Liability — check the PDF")
                    warnings.append(f"{fname}: table totals do not tie to Total Liability")
                info["note"] = "; ".join(notes)
                continue
            p3 = parse_gstr3b_outward_pdf(content)
            if p3:
                g, per = p3["gstin"], p3["period"]
                info.update(kind="GSTR-3B (portal PDF)", gstin=g, state=state_of_gstin(g), periods=[per])
                if per in g3b_pdf[g]:
                    info.update(status="Skipped", note=f"{per} already read from {g3b_pdf[g][per]['_file']}")
                    continue
                p3["_file"] = fname
                g3b_pdf[g][per] = p3
                info["status"] = "Used" if p3["parsed_ok"] else "Used (3.1 not found)"
                continue
            info.update(kind="PDF", status="Skipped", note="Not a GSTR-1 or GSTR-3B portal PDF")
            warnings.append(f"{fname}: not recognised as a GSTR-1 or GSTR-3B PDF — ignored")
            continue

        try:
            gstr1_df, gstr3b_df, gstr2b_df = read_octa_excel(item)
        except Exception as e:
            info.update(kind="Excel", status="Skipped", note=f"Could not read: {e}")
            warnings.append(f"{fname}: could not be read ({e})")
            continue
        info["kind"] = "GSTR-1 (OCTA Excel)" if not gstr1_df.empty else "OCTA Excel"
        if not gstr2b_df.empty:
            raw_2b.append(gstr2b_df.assign(**{"Source File": fname}))
        m = _GSTIN_RE.search(fname.upper())
        file_gstin = m.group(0) if m else ""
        if gstr1_df.empty:
            if gstr3b_df.empty:
                info.update(status="Skipped", note="No Final GSTR-1 or GSTR-3B sheet found")
                warnings.append(f"{fname}: no GSTR-1 / GSTR-3B sheet found — ignored")
            else:
                info.update(kind="GSTR-3B (OCTA Excel)", status="Used")
                pending_3b_sheets.append((fname, gstr3b_df, file_gstin))
            continue
        cg_col = _find_col(gstr1_df, ["Company GSTIN", "Registration GSTIN", "Supplier GSTIN"])
        row_g = gstr1_df[cg_col].map(_clean_gstin) if cg_col else pd.Series([""] * len(gstr1_df), index=gstr1_df.index)
        known = [g for g in row_g if g]
        fallback = Counter(known).most_common(1)[0][0] if known else file_gstin
        if not fallback:
            undetected.append((fname, gstr1_df, gstr3b_df, info))
            continue
        row_g = row_g.map(lambda g: g or fallback)
        octa_pending.append((fname, gstr1_df, row_g, info))
        if not gstr3b_df.empty:
            pending_3b_sheets.append((fname, gstr3b_df, fallback))

    # When two files cover the same registration x month (a yearly export plus the
    # monthly ones), the file covering FEWER periods wins — the monthly return is the
    # one filed for that month — so the result never depends on upload order.
    def _n_periods(item):
        df = item[1]
        pc = _find_col(df, ["Tax Period", "Tax period", "Period", "Return Period"])
        return len({p for p in df[pc].map(_period_key) if p}) if pc else 999
    for fname, gstr1_df, row_g, info in sorted(octa_pending, key=lambda it: (_n_periods(it), it[0])):
        _ingest_octa(gstr1_df, row_g, fname, info, octa_seen, octa_frames, raw_g1)

    for fname, gstr1_df, gstr3b_df, info in undetected:
        info.update(status="Skipped", note="No GSTIN in the file or its name — cannot tell which state it is")
        warnings.append(f"{fname}: no GSTIN in the file or its name — ignored. Put the GSTIN in the file name.")

    regs_known = sorted(set(octa_frames) | set(g1_pdf) | set(g3b_pdf))
    for fname, gdf, file_gstin in pending_3b_sheets:
        raw_3b.append(gdf.assign(**{"Source File": fname}))
        cg = _find_col(gdf, ["Company GSTIN", "GSTIN", "State"])
        keys = gdf[cg].fillna("").astype(str).str.strip() if cg else pd.Series([""] * len(gdf), index=gdf.index)
        for val in keys.unique():
            part = gdf[keys == val].reset_index(drop=True)
            reg = _state_text_to_reg(val, regs_known) if val else ""
            if not reg and file_gstin and (not val or _base_state(val) == _base_state(state_of_gstin(file_gstin))):
                reg = file_gstin
            if not reg:
                continue
            for month, amt in extract_gstr3b_monthly(part).items():
                if any(abs(amt[k]) > 0 for k in _AMOUNT_KEYS) and month not in g3b_octa[reg]:
                    g3b_octa[reg][month] = amt

    all_periods = [p for (_g, p) in octa_seen] + \
                  [p for g in g1_pdf for p in g1_pdf[g]] + [p for g in g3b_pdf for p in g3b_pdf[g]]
    fys = sorted({_fy_of(p) for p in all_periods if _fy_of(p)})
    if len(fys) > 1:
        warnings.append("Returns cover more than one financial year (" + ", ".join(fys) + "). Months are "
                        "compared by name (April…March), so upload one financial year per run.")

    return {
        "octa": {g: _period_sorted(pd.concat(fr, ignore_index=True)) for g, fr in octa_frames.items()},
        "g1_pdf": {g: dict(v) for g, v in g1_pdf.items()},
        "g3b_pdf": {g: dict(v) for g, v in g3b_pdf.items()},
        "g3b_octa": {g: dict(v) for g, v in g3b_octa.items()},
        "raw_gstr1": pd.concat(raw_g1, ignore_index=True) if raw_g1 else pd.DataFrame(),
        "raw_gstr3b": pd.concat(raw_3b, ignore_index=True) if raw_3b else pd.DataFrame(),
        "raw_gstr2b": pd.concat(raw_2b, ignore_index=True) if raw_2b else pd.DataFrame(),
        "files": files,
        "warnings": warnings,
        "fys": fys,
    }


def _period_sorted(df: pd.DataFrame) -> pd.DataFrame:
    """Monthly files arrive in any order (often by file name: Jan before Apr). Put
    rows back in tax-period order — stable, so a single file keeps its own order."""
    col = _find_col(df, ["Tax Period", "Tax period", "Period", "Return Period"])
    if col is None or df.empty:
        return df
    keys = df[col].map(_period_key)
    return df.iloc[sorted(range(len(df)), key=lambda i: (keys.iloc[i] == "", keys.iloc[i]))].reset_index(drop=True)


def _ingest_octa(gstr1_df, row_g, fname, info, seen, frames, raw_frames):
    """Keep this file's rows, minus any registration x period an earlier file gave."""
    period_col = _find_col(gstr1_df, ["Tax Period", "Tax period", "Period", "Return Period"])
    periods = gstr1_df[period_col].map(_period_key) if period_col else pd.Series([""] * len(gstr1_df), index=gstr1_df.index)
    keys = list(zip(row_g, periods))
    dropped = {k for k in set(keys) if k[1] and k in seen}
    keep = pd.Series([k not in dropped for k in keys], index=gstr1_df.index)
    for k in set(keys):
        if k[1] and k not in seen:
            seen[k] = fname
    kept, kept_g = gstr1_df[keep], row_g[keep]
    gst = sorted({g for g in row_g if g})
    info.update(gstin=", ".join(gst), state=", ".join(state_of_gstin(g) for g in gst),
                periods=sorted({p for (g, p) in keys if p and (g, p) not in dropped}),
                rows=int(keep.sum()), status="Used" if keep.any() else "Skipped")
    if dropped:
        info["note"] = (f"{int((~keep).sum())} row(s) skipped — period(s) "
                        + ", ".join(sorted({p for _g, p in dropped})) + " already read from an earlier file")
    for g in kept_g.unique():
        frames[g].append(kept[kept_g == g].reset_index(drop=True))
    raw = kept.copy()
    raw.insert(0, "Registration GSTIN", kept_g.values)
    raw.insert(0, "Source File", fname)
    raw_frames.append(raw)


# ---------------------------------------------------------------------------
# The existing single-state steps, run on one registration
# ---------------------------------------------------------------------------

def _gstr1_split_monthly(gstr1_df: pd.DataFrame) -> tuple[dict, dict]:
    """Same B2B/B2C monthly split the single-state server branch computes inline."""
    b2b, b2c = defaultdict(_zero_amounts), defaultdict(_zero_amounts)
    if gstr1_df is None or gstr1_df.empty:
        return {}, {}
    month_of, _ = gstr1_month_basis(gstr1_df)
    gstin_col = _find_col(gstr1_df, ["Customer GSTIN", "GSTIN of Recipient"])
    taxable_col = _find_col(gstr1_df, ["Item Taxable Value", "Taxable Value"])
    igst_col = _find_col(gstr1_df, ["IGST", "Integrated Tax"])
    cgst_col = _find_col(gstr1_df, ["CGST", "Central Tax"])
    sgst_col = _find_col(gstr1_df, ["SGST", "State Tax"])
    for _, row in gstr1_df.iterrows():
        month = month_of(row)
        if month not in _FY_MONTHS:
            continue
        gstin = str(_col_val(row, gstin_col, "")).strip()
        is_b2b = bool(gstin) and len(re.sub(r"[^A-Z0-9]", "", gstin.upper())) == 15
        _add_amounts(b2b[month] if is_b2b else b2c[month],
                     taxable=_f(_col_val(row, taxable_col, 0)), igst=_f(_col_val(row, igst_col, 0)),
                     cgst=_f(_col_val(row, cgst_col, 0)), sgst=_f(_col_val(row, sgst_col, 0)))
    return dict(b2b), dict(b2c)


def _fill_months(base: dict, extra: dict) -> dict:
    """Months already present from the invoice-level file win; the portal summary
    fills only the months that file does not cover."""
    out = {m: dict(v) for m, v in (base or {}).items()}
    for m, v in (extra or {}).items():
        if m not in out:
            out[m] = dict(v)
    return out


def run_registration(books_df, octa_df, pdf_periods: dict, g3b_monthly: dict,
                     tolerance: float, register_cols: list, invoice_mode: bool,
                     split_ok: bool) -> dict:
    octa_df = octa_df if octa_df is not None else pd.DataFrame()
    octa_monthly = aggregate_gstr1_monthly(octa_df)
    octa_b2b, octa_b2c = _gstr1_split_monthly(octa_df)
    pdf_all = {v["month"]: _amt(v["total"]) for v in (pdf_periods or {}).values()}
    pdf_b2b = {v["month"]: _amt(v["b2b"]) for v in (pdf_periods or {}).values()}
    pdf_b2c = {v["month"]: _amt(v["b2c"]) for v in (pdf_periods or {}).values()}

    gstr1_monthly = _fill_months(octa_monthly, pdf_all)
    g1_b2b_m = _fill_months(octa_b2b, pdf_b2b)
    g1_b2c_m = _fill_months(octa_b2c, pdf_b2c)

    g1_b2b_invs = _gstr1_b2b_invoice_set(octa_df)
    books_all = aggregate_books_monthly(books_df)
    books_b2b = aggregate_books_monthly(books_df, category="B2B", g1_b2b_invs=g1_b2b_invs)
    books_b2c = aggregate_books_monthly(books_df, category="B2C", g1_b2b_invs=g1_b2b_invs)

    # A GSTR-3B comparison only means something for the months a 3B was uploaded;
    # every other month would read as the whole turnover "missing" from 3B.
    g3b_months = set(g3b_monthly or {})
    books_3b = {m: v for m, v in books_all.items() if m in g3b_months}
    gstr1_3b = {m: v for m, v in gstr1_monthly.items() if m in g3b_months}
    sections = {
        "gstr1_vs_gstr3b": build_monthly_comparison(gstr1_3b, g3b_monthly or {}, "gstr1", "gstr3b"),
        "books_all_vs_gstr1": build_monthly_comparison(books_all, gstr1_monthly, "books", "gstr1"),
        "books_b2b_vs_gstr1": build_monthly_comparison(books_b2b, g1_b2b_m, "books", "gstr1"),
        "books_b2c_vs_gstr1": build_monthly_comparison(books_b2c, g1_b2c_m, "books", "gstr1"),
        "books_all_vs_gstr3b": build_monthly_comparison(books_3b, g3b_monthly or {}, "books", "gstr3b"),
    }

    b2b_rows, b2c_rows = [], []
    if invoice_mode and not octa_df.empty:
        # reconcile_b2b_new returns nothing for an empty register, which would hide
        # this registration's GSTR-1 invoices. One blank row (skipped by it) keeps
        # every GSTR-1 line in the output as "Not in Books".
        b2b_books = books_df if not books_df.empty else pd.DataFrame([{c: None for c in register_cols}])
        b2b_rows = reconcile_b2b_new(b2b_books, octa_df, tolerance)
        b2c_books = books_df if not books_df.empty else pd.DataFrame(columns=register_cols)
        b2c_rows = reconcile_b2c_new(b2c_books, octa_df, tolerance)
    elif invoice_mode and not books_df.empty:
        b2b_rows = reconcile_b2b_new(books_df, octa_df, tolerance)

    pivot = None
    if not octa_df.empty and pdf_all:
        both = {m: v for m, v in pdf_all.items() if m in octa_monthly}
        if both:
            pivot = build_monthly_comparison({m: octa_monthly[m] for m in both}, both, "excel", "pdf")

    return {
        "sections": sections,
        "b2b_rows": b2b_rows,
        "b2c_rows": b2c_rows,
        "pivot_rows": pivot,
        "gstr3b_available": bool(g3b_monthly),
        "split_ok": split_ok,
        "summary": build_summary(b2b_rows, b2c_rows) if (b2b_rows or b2c_rows) else {},
        "totals": {"books": _sum_months(books_all), "gstr1": _sum_months(gstr1_monthly),
                   "gstr3b": _sum_months(g3b_monthly or {})},
        "_monthly": {"books": books_all, "gstr1": gstr1_monthly, "gstr3b": g3b_monthly or {}},
    }


# ---------------------------------------------------------------------------
# Orchestrator
# ---------------------------------------------------------------------------

def run_gstr1_multistate(tally_df: pd.DataFrame, cn_df: pd.DataFrame | None,
                         return_items: list[dict], tolerance: float,
                         tally_files: list[dict] | None = None) -> dict:
    ret = read_return_files(return_items)
    regs = sorted(set(ret["octa"]) | set(ret["g1_pdf"]) | set(ret["g3b_pdf"]) | set(ret["g3b_octa"]),
                  key=lambda g: (state_of_gstin(g) or g, g))
    if not regs:
        raise ValueError("None of the uploaded return files could be read as a GSTR-1 or GSTR-3B "
                         "(OCTA Excel or GST portal PDF).")

    reco_df = merge_credit_notes(tally_df, cn_df)
    reco_df, returns_info = normalise_returns(reco_df)
    reco_df = reco_df.reset_index(drop=True)
    register_cols = list(reco_df.columns)
    reg_of, reason_of = assign_books_states(reco_df, ret["octa"], regs)
    invoice_mode = bool(ret["octa"])
    split_ok = books_split_available(reco_df)

    g3b_by_reg: dict[str, dict] = {}
    for reg in regs:
        m3 = {m: _amt(v) for m, v in (ret["g3b_octa"].get(reg) or {}).items()}
        for per, p3 in (ret["g3b_pdf"].get(reg) or {}).items():
            month = p3["month"]
            if month not in m3:
                m3[month] = _amt(p3["outward"])
        g3b_by_reg[reg] = m3

    noreg = sorted({r for r in reg_of if str(r).startswith(NOREG)})
    groups = regs + noreg + ([UNASSIGNED] if (reg_of == UNASSIGNED).any() else [])
    states: list[dict] = []
    for reg in groups:
        books_slice = reco_df[reg_of == reg].reset_index(drop=True)
        res = run_registration(books_slice, ret["octa"].get(reg), ret["g1_pdf"].get(reg, {}),
                               g3b_by_reg.get(reg, {}), tolerance, register_cols, invoice_mode, split_ok)
        label = reg_label(reg)
        gstin = reg if is_real_reg(reg) else ""
        for r in res["b2b_rows"]:
            r["_reg_state"], r["_reg_gstin"] = label, gstin
        for r in res["b2c_rows"]:
            r["registration"] = label
        res.update({"state": label, "gstin": gstin, "books_rows": int(len(books_slice)),
                    "gstr1_rows": int(len(ret["octa"].get(reg, pd.DataFrame()))),
                    "gstr1_source": _source_of(ret, reg)})
        states.append(res)

    if invoice_mode:
        _cross_state_remarks(states)

    all_sections = {}
    for key, (lp, rp) in {"gstr1_vs_gstr3b": ("gstr1", "gstr3b"), "books_all_vs_gstr1": ("books", "gstr1"),
                           "books_b2b_vs_gstr1": ("books", "gstr1"), "books_b2c_vs_gstr1": ("books", "gstr1"),
                           "books_all_vs_gstr3b": ("books", "gstr3b")}.items():
        left, right = defaultdict(_zero_amounts), defaultdict(_zero_amounts)
        for st in states:
            for row in st["sections"].get(key) or []:
                if row.get("month") == "Total":
                    continue
                for k in _AMOUNT_KEYS:
                    left[row["month"]][k] += row.get(f"{lp}_{k}", 0) or 0
                    right[row["month"]][k] += row.get(f"{rp}_{k}", 0) or 0
        all_sections[key] = build_monthly_comparison(dict(left), dict(right), lp, rp)

    all_b2c = []
    if invoice_mode:
        all_octa = pd.concat(list(ret["octa"].values()), ignore_index=True)
        all_b2c = reconcile_b2c_new(reco_df, all_octa, tolerance)

    unassigned_df = reco_df[reg_of == UNASSIGNED].copy()
    unassigned_df.insert(0, "Why Unassigned", reason_of[reg_of == UNASSIGNED].values)
    books_with_reg = reco_df.copy()
    books_with_reg.insert(0, "Assigned By", reason_of.values)
    books_with_reg.insert(0, "Assigned Registration", [reg_display(r) for r in reg_of])

    grand = None
    if tally_files:
        parts = [register_grand_total(f) for f in tally_files]
        if all(p is not None for p in parts):
            grand = {k: round_money(sum(p[k] for p in parts)) for k in _AMOUNT_KEYS}
    row_checks = register_row_checks(reco_df)
    if row_checks:
        warnings_rows = len(row_checks)
    else:
        warnings_rows = 0
    checks = _integrity_checks(reco_df, states, ret, grand, cn_df is not None and not cn_df.empty)
    if row_checks:
        for c in checks:
            if not c["ok"] and "Grand Total" in c["check"]:
                c["note"] = ("Gap of " + format(round_money(abs(c["expected"] - c["actual"])), ",.2f")
                             + " — see 'Register Row Checks' for the row(s) that do not foot")
    status_grid = _return_status(states, ret, regs)

    state_summary = []
    for st in states:
        t = st["totals"]
        state_summary.append({
            "state": st["state"], "gstin": st["gstin"], "gstr1_source": st["gstr1_source"],
            "books_rows": st["books_rows"], "gstr1_rows": st["gstr1_rows"],
            **{f"books_{k}": t["books"][k] for k in _AMOUNT_KEYS},
            **{f"gstr1_{k}": t["gstr1"][k] for k in _AMOUNT_KEYS},
            **{f"gstr3b_{k}": t["gstr3b"][k] for k in _AMOUNT_KEYS},
            **{f"diff_books_gstr1_{k}": round_money(t["books"][k] - t["gstr1"][k]) for k in _AMOUNT_KEYS},
            **{f"diff_gstr1_gstr3b_{k}": round_money(t["gstr1"][k] - t["gstr3b"][k]) for k in _AMOUNT_KEYS},
            "gstr3b_available": st["gstr3b_available"],
            "b2b_summary": st["summary"],
        })

    warnings = list(ret["warnings"])
    n_un = int((reg_of == UNASSIGNED).sum())
    for grp in noreg:
        n = int((reg_of == grp).sum())
        warnings.append(f"{n} Sales Register row(s) belong to {grp[len(NOREG):]}, but no return for "
                        f"{grp[len(NOREG):]} was uploaded — shown separately, not compared.")
    if n_un:
        warnings.append(f"{n_un} Sales Register row(s) could not be tied to a registration — see the "
                        "'Unassigned Books' sheet. They are still counted in All States.")
    if not split_ok:
        warnings.append("The Sales Register has no Category or buyer-GSTIN column, so Books cannot be split "
                        "into B2B and B2C — compare 'All sales'. The B2B/B2C sections are left out.")
    if warnings_rows:
        warnings.append(f"{warnings_rows} Sales Register row(s) whose ledger columns do not add up to their "
                        "Gross Total — usually one cell exported with the wrong sign. Listed in "
                        "'Register Row Checks'; the amounts are used as exported.")
    if returns_info["flipped"]:
        warnings.append(f"{returns_info['flipped']} row(s) marked as returns in '{returns_info['column']}' "
                        "carried positive amounts — they were netted off (sign-flipped) against sales.")

    return {
        "states": states,
        "state_summary": state_summary,
        "all_sections": all_sections,
        "all_b2b": [r for st in states for r in st["b2b_rows"]],
        "all_b2c": all_b2c,
        "b2c_by_reg": [r for st in states for r in st["b2c_rows"]],
        "unassigned_df": unassigned_df,
        "books_with_reg": books_with_reg,
        "checks": checks,
        "status_grid": status_grid,
        "files": ret["files"],
        "warnings": warnings,
        "raw_gstr1": ret["raw_gstr1"],
        "raw_gstr3b": ret["raw_gstr3b"],
        "raw_gstr2b": ret["raw_gstr2b"],
        "g1_pdf": ret["g1_pdf"],
        "g3b_pdf": ret["g3b_pdf"],
        "register_cols": register_cols,
        "invoice_mode": invoice_mode,
        "split_ok": split_ok,
        "returns_info": returns_info,
        "unassigned_rows": n_un,
        "regs": regs,
        "register_row_checks": row_checks,
        "grand_total": grand,
    }


def _source_of(ret: dict, reg: str) -> str:
    parts = []
    if reg in ret["octa"]:
        parts.append("OCTA Excel")
    if ret["g1_pdf"].get(reg):
        parts.append(f"Portal PDF x{len(ret['g1_pdf'][reg])}")
    return " + ".join(parts) or ("—" if reg != UNASSIGNED else "")


def _return_status(states: list[dict], ret: dict, regs: list[str]) -> list[dict]:
    """Registration x month: which return was read, and does Books have turnover there."""
    rows = []
    for st in states:
        reg = st["gstin"]
        if not reg:
            continue
        octa = ret["octa"].get(reg)
        octa_months = set(aggregate_gstr1_monthly(octa).keys()) if octa is not None else set()
        pdfs = {v["month"]: v for v in (ret["g1_pdf"].get(reg) or {}).values()}
        pdf3 = {v["month"] for v in (ret["g3b_pdf"].get(reg) or {}).values()}
        octa3 = set((ret["g3b_octa"].get(reg) or {}).keys())
        books_m = st["_monthly"]["books"]
        for m in _FY_MONTHS:
            has_books = any(abs(_f((books_m.get(m) or {}).get(k, 0))) > 0 for k in _AMOUNT_KEYS)
            if m in octa_months:
                g1 = "OCTA Excel" + (" + PDF" if m in pdfs else "")
            elif m in pdfs:
                g1 = "Portal PDF (Nil return)" if pdfs[m]["nil_filed"] else "Portal PDF"
            else:
                g1 = "Not uploaded"
            g3 = "Portal PDF" if m in pdf3 else ("OCTA Excel" if m in octa3 else "Not uploaded")
            flag = ""
            if has_books and g1 == "Not uploaded":
                flag = "Books has sales but no GSTR-1 uploaded"
            elif has_books and "Nil" in g1:
                flag = "Books has sales but GSTR-1 was filed Nil"
            rows.append({"state": st["state"], "gstin": reg, "month": m, "books_has_sales": has_books,
                         "gstr1": g1, "gstr3b": g3, "flag": flag})
    return rows


_INV_HEADER_KEYS = ("voucherno", "voucherrefno", "invoiceno", "docno", "billno")


def _cross_state_remarks(states: list[dict]) -> None:
    """Remark 3: booked under one registration, filed in another's GSTR-1."""
    not_in_books: dict[str, list] = defaultdict(list)
    for st in states:
        for r in st["b2b_rows"]:
            r.setdefault("_remark3", "")
            if r.get("_remark") in ("Not in Books", "Amazon Entry as per GSTR-1"):
                n = _norm_inv(r.get("_gstr1_inv_no"))
                if n:
                    not_in_books[n].append((st, r))
    for st in states:
        inv_keys = None
        for r in st["b2b_rows"]:
            if r.get("_remark") != "Not in GSTR-1":
                continue
            if inv_keys is None:
                inv_keys = [c for c in r if not str(c).startswith("_")
                            and re.sub(r"[^a-z]", "", str(c).lower()) in _INV_HEADER_KEYS]
            hit = None
            for c in inv_keys:
                n = _norm_inv(r.get(c))
                for other, g_row in (not_in_books.get(n, []) if n else []):
                    if other is not st:
                        hit = (other, g_row)
                        break
                if hit:
                    break
            if not hit:
                continue
            other, g_row = hit
            here = st["state"]
            there = f"{other['state']} ({other['gstin']})" if other["gstin"] else other["state"]
            r["_remark3"] = (f"Filed in {there} GSTR-1 — this row belongs to {other['state']}" if here == UNASSIGNED
                             else f"Filed in {there} GSTR-1, but booked under {here} in Tally — check the voucher type")
            g_row["_remark3"] = (f"Booked in Tally under {here}" if here != UNASSIGNED
                                 else "In Tally, but the row's state could not be identified")


def _integrity_checks(reco_df, states, ret, grand, has_cn) -> list[dict]:
    """Prove nothing was dropped or double counted."""
    checks = []
    n_rows, n_split = len(reco_df), sum(st["books_rows"] for st in states)
    checks.append({"check": "Every Sales Register row sits in exactly one registration (incl. Unassigned)",
                   "expected": n_rows, "actual": n_split, "ok": n_rows == n_split})

    whole = _sum_months(aggregate_books_monthly(reco_df))
    split = {k: round_money(sum(st["totals"]["books"][k] for st in states)) for k in _AMOUNT_KEYS}
    checks.append({"check": "Books value — sum of registrations = whole register (taxable)",
                   "expected": whole["taxable"], "actual": split["taxable"],
                   "ok": all(abs(whole[k] - split[k]) < 0.05 for k in _AMOUNT_KEYS)})
    if grand is not None and not has_cn:
        checks.append({"check": "Books taxable ties to Tally's Grand Total row",
                       "expected": grand["taxable"], "actual": whole["taxable"],
                       "ok": abs(grand["taxable"] - whole["taxable"]) < 1.0})
        tax_g = round_money(grand["igst"] + grand["cgst"] + grand["sgst"])
        tax_b = round_money(whole["igst"] + whole["cgst"] + whole["sgst"])
        checks.append({"check": "Books tax (IGST+CGST+SGST) ties to Tally's Grand Total row",
                       "expected": tax_g, "actual": tax_b, "ok": abs(tax_g - tax_b) < 1.0})

    g_rows = sum(len(df) for df in ret["octa"].values())
    if g_rows:
        g_split = sum(st["gstr1_rows"] for st in states)
        checks.append({"check": "Every GSTR-1 (OCTA) row sits in exactly one registration",
                       "expected": g_rows, "actual": g_split, "ok": g_rows == g_split})
        for st in states:
            g1 = ret["octa"].get(st["gstin"]) if st["gstin"] else None
            if g1 is None or g1.empty:
                continue
            gc = _find_col(g1, ["Customer GSTIN", "GSTIN of Recipient", "Buyer GSTIN"])
            n_b2b = 0
            if gc:
                for _, row in g1.iterrows():
                    v = str(_col_val(row, gc, "")).strip()
                    if v and v.upper() not in ("N/A", "NA"):
                        n_b2b += 1
            n_in = sum(1 for r in st["b2b_rows"] if str(r.get("_gstr1_gstin") or "").strip())
            checks.append({"check": f"{st['state']}: every GSTR-1 B2B line appears in the B2B Reco",
                           "expected": n_b2b, "actual": n_in, "ok": n_in == n_b2b})

    pdfs = [p for g in ret["g1_pdf"].values() for p in g.values()]
    if pdfs:
        bad = [p["_file"] for p in pdfs if not p["parsed_ok"]]
        checks.append({"check": "Every GSTR-1 PDF: table totals = the form's own Total Liability line",
                       "expected": len(pdfs), "actual": len(pdfs) - len(bad), "ok": not bad})
        pdf_sum = round_money(sum(_f(p["total"]["taxable"]) for p in pdfs))
        used = round_money(sum(st["totals"]["gstr1"]["taxable"] for st in states
                               if st["gstin"] and st["gstin"] not in ret["octa"]))
        pdf_only = round_money(sum(_f(p["total"]["taxable"]) for g, per in ret["g1_pdf"].items()
                                   if g not in ret["octa"] for p in per.values()))
        checks.append({"check": "GSTR-1 PDFs — every PDF's taxable value is in the reco",
                       "expected": pdf_only, "actual": used, "ok": abs(pdf_only - used) < 1.0})
    p3 = [p for g in ret["g3b_pdf"].values() for p in g.values()]
    if p3:
        bad = [p["_file"] for p in p3 if not p["parsed_ok"]]
        checks.append({"check": "Every GSTR-3B PDF: table 3.1(a) found",
                       "expected": len(p3), "actual": len(p3) - len(bad), "ok": not bad})
    return checks


def records(df: pd.DataFrame | None) -> list[dict]:
    return df_to_records(df) if df is not None and not df.empty else []
